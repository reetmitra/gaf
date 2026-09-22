"""The golden-set harness: concurrent validation, and the pipeline regression fixture.

This module is the project's safety net. It exists so that a prompt edit, a threshold
change, a mock-persona rewrite or an accidental reordering becomes a **regression-testable
event** rather than something a reader notices three weeks later in the write-up.

It has two jobs, and they are deliberately kept apart.

--------------------------------------------------------------------------------
JOB (a) — concurrent validation against a human coding
--------------------------------------------------------------------------------

Alqazlan-style: `gaf.analysis.agreement.concurrent_validation` is called, never
reimplemented. The harness's own contribution is the **loading** of a human coding in
the row-oriented assignments shape ``[{"response_id": int, "segment": str, "code": str}]``
— the shape of the PI's coding spreadsheets, and exactly what
`gaf.ingest.xlsx.read_coded_xlsx` produces from one.

**The PI's real hand-codings do not exist in this repository and never will.** What sits
in `tests/fixtures/golden/human_coding.json` is a synthetic stand-in of identical shape.

*Swapping the real file in — the single path, no code change:*

    export GAF_GOLDEN_HUMAN_CODING="../Grounded AI Futures/data/IndiaHandCoding.xlsx"

`load_human_coding()` then reads that path instead of the synthetic stand-in, dispatching
on the suffix: ``.xlsx``/``.xlsm`` through `read_coded_xlsx`, ``.json`` through
`Assignment.from_json`. `validate_against_golden()` is the harness entry point and takes
the same override. The file stays outside the repository; nothing is copied in.

The numeric assertions below deliberately keep reading the *synthetic* stand-in, because
they are the harness's own self-test. `test_real_golden_set_loads_when_supplied` is the
one test that follows the override, and it is skipped whenever the variable is unset.

*The arithmetic, so a human can check it by hand.* Re-derived by hand against the 200-range
synthetic fixture (ADR-0024); the cell *pattern* was preserved by the rewrite, so every
number below is unchanged from the pre-rewrite derivation except the response ids.

The synthetic golden coding has 8 rows over responses 203, 207, 226, 239, 244, 245; the
machine coding (`tests.fixtures.assignments`) has 8 rows over responses 203, 207, 212,
226, 244, 245, 258 — 8 distinct responses between them. Four code names are shared
verbatim and so match at cosine 1.0 (`positive_impacts-healthcare`,
`negative_impacts-job_destruction`, `positive_impacts-problem-solving`,
`negative_impacts-misuse`); `future-inevitability` (machine-only) and
`negative_impacts-lack_of_inclusion` (human-only) match nothing above tau_high = 0.80,
so the union vocabulary is 4 + 1 + 1 = 6 columns. The response x code matrix is therefore
8 responses x 6 columns = 48 cells:

    TP = 6   (203,healthcare) (203,job_destruction) (207,job_destruction)
             (226,healthcare) (244,problem-solving) (245,misuse)
    FP = 2   (212,inevitability) (258,inevitability)
    FN = 2   (207,lack_of_inclusion) (239,lack_of_inclusion)
    TN = 48 - 6 - 2 - 2 = 38

    precision = 6/(6+2) = 0.75   recall = 6/(6+2) = 0.75   F1 = 0.75

    observed  = (6 + 38)/48 = 44/48 = 0.916666...
    expected  = (8/48)(8/48) + (40/48)(40/48) = (64 + 1600)/2304
              = 1664/2304 = 0.722222...
    kappa     = (2112/2304 - 1664/2304) / (2304/2304 - 1664/2304)
              = 448/640 = 0.7 exactly

On the 4 matched columns alone the machine is perfect — the same 6 cells are present on
both sides, so TP = 6, FP = FN = 0, F1 = 1.0 and kappa = 1.0.

Segment level catches what code level cannot. Every one of the 8 human and 8 machine
segments locates in the corpus, and each row is scored by the best span-Jaccard against a
row of the other coding carrying the same code on the same response:

    machine rows: five exact-span matches at 1.0
                  (203,job_destruction) "no retraining scheme arrives" [151,179)
                      vs human [93,145)              -> disjoint, 0.0
                  (212,inevitability), (258,inevitability) -> no human row, 0.0
    human rows:   the same five at 1.0
                  (203,job_destruction) [93,145)     -> 0.0
                  (207,lack_of_inclusion), (239,lack_of_inclusion) -> 0.0

    weighted precision = 5/8 = 0.625    weighted recall = 5/8 = 0.625
    weighted F1        = 2(0.625)(0.625)/1.25 = 0.625

The machine's second row on response 203 carries the right code on the wrong span, so
`right_code_wrong_place` is 1 and the weighted scores fall to 0.625 while the code-level
cell for (203, job_destruction) still counts as a hit. That row is the whole reason the
two levels are reported separately.

--------------------------------------------------------------------------------
JOB (b) — pipeline regression
--------------------------------------------------------------------------------

The offline fast loop is run over `tests.fixtures.corpus.synthetic_corpus()` with mock
clients, and its output is compared byte-for-byte against artefacts committed under
`tests/fixtures/golden/`:

    codebook.json      the canonical `FastLoopResult.codebook_json()` bytes
    assignments.json   every assignment row, in the loop's own total order
    snapshots.json     the snapshot-id sequence (content hashes of the codebook)
    findings.json      the check x severity summary, and every ERROR the run expects
    manifest.json      the run configuration, thresholds, embedding space and prompt
                       digests the fixture was generated under

`manifest.json` is what turns "a threshold moved" or "a prompt was edited" from a silent
behaviour change into a named line in a failure message: every `CodingRules` field and a
digest of every prompt template's static text are pinned there, so an edit fails even in
the cases where the mock clients would have absorbed it.

*Regenerating the fixture, on purpose:*

    uv run python -m tests.test_golden --regenerate

The flag is required — running the module without it exits with a usage error — so the
fixture cannot be rewritten by accident, and no environment variable can do it either.
`human_coding.json` is **not** regenerated: it is the slot the PI's file replaces, and a
regeneration must never overwrite a human coding.

Beyond the bytes, four properties are pinned that must never regress:

* determinism — two runs are byte-identical (acceptance criterion 3);
* order independence — shuffled corpora produce the same final codebook (ADR-0021,
  verified there on the real data; it must stay true here);
* the fast loop never emits `split`, `reparent` or `rename` (ADR-0021, and the
  documented cause of the predecessor's codebook flattening);
* no ERROR-severity finding appears that `findings.json` does not expect.

--------------------------------------------------------------------------------
Real-data safety
--------------------------------------------------------------------------------

Human survey responses must never enter git history. Every fixture here is invented, and
`test_no_real_survey_data_is_committed` proves it from the filesystem rather than from a
promise: no file with a real-data extension exists anywhere the repository tracks, the
`.gitignore` guards are still in place, every file under `tests/fixtures/golden/` is
JSON, and **every segment in every golden fixture is locatable in the synthetic corpus**
— which is what makes "nothing here derives from a respondent's words" checkable rather
than asserted.

Validation principles: **reliability** (the regression fixture is what makes byte-identity
a claim with a test behind it) and **interpretive depth** (the over-coding and blind-spot
lists, not the headline percentage, are what the write-up quotes).
"""

from __future__ import annotations

import argparse
import difflib
import importlib.util
import json
import os
import random
import re
import sys
from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest

from gaf.agents.prompts.loader import LOOP_ROLES, all_templates
from gaf.analysis.agreement import AgreementReport, concurrent_validation
from gaf.config import QUESTION_V2, CodingRules, EmbeddingSpaceConfig, RunConfig
from gaf.embed.protocol import Embedder
from gaf.embed.service import EmbeddingService
from gaf.ids import content_hash
from gaf.ingest.xlsx import read_coded_xlsx
from gaf.models import OPERATION_TYPES, Assignment, Code, Codebook, Evidence, Response
from gaf.pipeline.fast_loop import FastLoopResult, offline_components, run_fast_loop
from gaf.store.blackboard import Blackboard
from gaf.textnorm import normalise
from tests.fixtures.assignments import (
    HUMAN_ONLY,
    MACHINE_ONLY,
    human_assignments,
    machine_assignments,
)
from tests.fixtures.corpus import corpus_by_id, response_ids, synthetic_corpus
from tests.fixtures.xlsx import write_coded_xlsx

# --------------------------------------------------------------------------- #
# Where the fixture lives
# --------------------------------------------------------------------------- #

TESTS_DIR = Path(__file__).resolve().parent
REPO_ROOT = TESTS_DIR.parent
GOLDEN_DIR = TESTS_DIR / "fixtures" / "golden"

CODEBOOK_FIXTURE = GOLDEN_DIR / "codebook.json"
ASSIGNMENTS_FIXTURE = GOLDEN_DIR / "assignments.json"
SNAPSHOTS_FIXTURE = GOLDEN_DIR / "snapshots.json"
FINDINGS_FIXTURE = GOLDEN_DIR / "findings.json"
MANIFEST_FIXTURE = GOLDEN_DIR / "manifest.json"

#: The drop-in slot for the PI's hand-coding. Never rewritten by `--regenerate`.
HUMAN_CODING_FIXTURE = GOLDEN_DIR / "human_coding.json"

#: Point this at the PI's real spreadsheet, outside the repository, and the harness
#: reads it with no code change. Documented in full in the module docstring.
HUMAN_CODING_ENV = "GAF_GOLDEN_HUMAN_CODING"

#: Printed in every drift message. The only way to rewrite the expected artefacts.
REGENERATE_COMMAND = "uv run python -m tests.test_golden --regenerate"

#: The run id the fixture was generated under. Nothing downstream depends on it —
#: ids are content-addressed (ADR-0018) — but the manifest records it anyway.
GOLDEN_RUN_ID = "golden"

#: Cap on the unified diff a failure prints. A full codebook diff is unreadable; the
#: headline above it already names what moved.
MAX_DIFF_LINES = 60

#: The suffixes `.gitignore` excludes repository-wide because a file carrying one is
#: almost certainly human survey data. Kept in step with the ignore file by
#: `test_no_real_survey_data_is_committed`.
REAL_DATA_SUFFIXES: frozenset[str] = frozenset({".xlsx", ".xlsm", ".docx", ".csv", ".tsv"})

#: Directory names `.gitignore` excludes, plus the VCS directory itself. Walking the
#: tree without pruning these would read scikit-learn's own bundled CSVs out of `.venv`
#: and report them as a data leak.
IGNORED_DIRS: frozenset[str] = frozenset(
    {
        ".git",
        ".venv",
        "venv",
        "__pycache__",
        ".pytest_cache",
        ".mypy_cache",
        ".ruff_cache",
        ".gaf_cache",
        ".idea",
        ".vscode",
        "htmlcov",
        "build",
        "dist",
        "node_modules",
        "runs",
    }
)

#: Response ids that appear in the *real* India seed sample and in no synthetic fixture
#: (ADR-0015 cites 18 and 56; ADR-0016 cites 9, 10 and 57; the coding rules cite 67).
#: Their presence in a golden fixture would mean real data had been copied in.
REAL_SAMPLE_ONLY_IDS: frozenset[int] = frozenset({9, 10, 18, 56, 57, 67})


# --------------------------------------------------------------------------- #
# Job (a) — loading a human coding, and running concurrent validation on it
# --------------------------------------------------------------------------- #


def golden_human_coding_path() -> Path:
    """The human coding the harness reads: the env override, or the synthetic stand-in."""
    override = os.environ.get(HUMAN_CODING_ENV)
    return Path(override).expanduser() if override else HUMAN_CODING_FIXTURE


def load_human_coding(path: Path | str | None = None) -> list[Assignment]:
    """Load a human coding in the row-oriented assignments shape.

    Dispatches on the suffix so the PI's workbook and the synthetic stand-in load through
    one call: ``.xlsx``/``.xlsm`` go through `gaf.ingest.xlsx.read_coded_xlsx` (which also
    repairs mojibake, so a human quote stays byte-comparable with the response it was
    taken from), and ``.json`` is read as a bare list of assignment rows or as
    ``{"assignments": [...]}``.
    """
    resolved = Path(path).expanduser() if path is not None else golden_human_coding_path()
    if not resolved.exists():
        raise FileNotFoundError(
            f"no human coding at {resolved}. Either the synthetic stand-in "
            f"{HUMAN_CODING_FIXTURE.name} is missing, or {HUMAN_CODING_ENV} points at a "
            "file that is not there."
        )
    if resolved.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_coded_xlsx(resolved)
    payload = json.loads(resolved.read_text(encoding="utf-8"))
    rows = payload["assignments"] if isinstance(payload, dict) else payload
    return [Assignment.from_json(row) for row in rows]


def golden_embedder() -> Embedder:
    """The production offline embedding space — the same one the fast loop runs in.

    Not `tests.fixtures.embedding.StubEmbedder`: a similarity reported by the agreement
    module has to be comparable with a similarity reported anywhere else in the system
    (ADR-0003), and that is only true inside one space.
    """
    return EmbeddingService(EmbeddingSpaceConfig())


def validate_against_golden(
    machine: Sequence[Assignment],
    *,
    human: Sequence[Assignment] | None = None,
    human_path: Path | str | None = None,
    corpus: Mapping[int, Response] | None = None,
    descriptions: Mapping[str, str] | None = None,
    rules: CodingRules | None = None,
    embedder: Embedder | None = None,
) -> AgreementReport:
    """The harness entry point: a machine coding against the human golden set.

    `human` wins if given; otherwise the coding is loaded from `human_path`, or from
    whatever `golden_human_coding_path()` resolves to. `corpus` is what makes the segment
    level a span comparison rather than string containment, so pass it whenever the
    responses are available.
    """
    return concurrent_validation(
        list(human) if human is not None else load_human_coding(human_path),
        list(machine),
        embedder=embedder or golden_embedder(),
        rules=rules or CodingRules(),
        corpus=corpus,
        descriptions=descriptions,
    )


# --------------------------------------------------------------------------- #
# Job (b) — the reference run and the artefacts it must reproduce
# --------------------------------------------------------------------------- #


def _run_prompt_hashes() -> dict[str, str]:
    """``role/version -> hash`` for the prompts a *run* consults.

    The Definer is registered in `gaf.agents.prompts.loader` with the other three but
    runs outside both loops, so it is excluded here: this manifest answers "what wording
    produced this coding", and the Definer produced none of it (ADR-0034, R1 N7a).
    """
    return {
        f"{template.role}/{template.version}": content_hash(template.all_text())
        for template in all_templates()
        if template.role in LOOP_ROLES
    }


def golden_config(run_id: str = GOLDEN_RUN_ID) -> RunConfig:
    """The configuration the committed fixture was generated under.

    `cache_dir=None` because a run that reads a disk cache is a run whose output depends
    on a directory outside the repository, which is the opposite of what this fixture is
    for.
    """
    return RunConfig(run_id=run_id, offline=True, cache_dir=None, batch_size=10)


def run_golden(
    run_id: str = GOLDEN_RUN_ID,
    responses: Sequence[Response] | None = None,
) -> tuple[FastLoopResult, Blackboard]:
    """One offline fast-loop run over the synthetic corpus, on an in-memory blackboard."""
    config = golden_config(run_id)
    board = Blackboard()
    result = run_fast_loop(
        list(responses if responses is not None else synthetic_corpus()),
        config=config,
        board=board,
        components=offline_components(config),
    )
    return result, board


def _dump(obj: Any) -> str:
    """The one serialisation every non-codebook fixture file uses."""
    return json.dumps(obj, indent=2, sort_keys=True, ensure_ascii=False) + "\n"


def expected_findings(result: FastLoopResult) -> dict[str, Any]:
    """The findings fixture: the summary, the totals, and every ERROR by identity.

    ERRORs are recorded by ``(check, scope, subject, marker)`` with a count rather than
    by message, so rewording a check's sentence does not fail the regression while a *new*
    ERROR — the thing that actually matters — does.
    """
    errors: dict[tuple[str, str, str, str], int] = {}
    for finding in result.report.errors():
        key = (
            finding.check_id,
            finding.scope,
            finding.subject,
            str(finding.data.get("marker", "")),
        )
        errors[key] = errors.get(key, 0) + 1
    return {
        "summary": result.report.summary(),
        "severities": result.report.totals(),
        "errors": [
            {
                "check_id": check_id,
                "scope": scope,
                "subject": subject,
                "marker": marker,
                "count": count,
            }
            for (check_id, scope, subject, marker), count in sorted(errors.items())
        ],
    }


def expected_manifest(result: FastLoopResult, config: RunConfig) -> dict[str, Any]:
    """Everything outside the artefacts that the artefacts depend on.

    A prompt edit that the mock clients happen to absorb, or a threshold moved by a
    fraction, would otherwise change nothing visible. Pinning the rules and a digest of
    each prompt template's static text makes both fail here, by name.
    """
    return {
        "regenerate_with": REGENERATE_COMMAND,
        "corpus": "tests.fixtures.corpus.synthetic_corpus",
        "response_ids": response_ids(),
        "run": {
            "run_id": config.run_id,
            "offline": config.offline,
            "seed": config.seed,
            "batch_size": config.batch_size,
            "snapshot_policy": config.snapshot_policy,
            "retrieval_top_k": config.retrieval_top_k,
            "question_variant": config.question_variant,
            "recycle_human_corrections": config.recycle_human_corrections,
        },
        "space_id": result.stats.space_id,
        "embedding": config.embedding.to_json(),
        "rules": config.rules.to_json(),
        "models": config.models.to_json(),
        # `LOOP_ROLES`, not every template in the build: this manifest records the
        # wording that produced *this coding*, and the Definer (registered in the
        # loader since R1 N7a) runs outside both loops and never sees a response being
        # coded. Its hash here would claim the coding depended on wording it never
        # read, and a Definer prompt edit would fail this fixture with that claim as
        # its message (ADR-0034).
        "prompts": _run_prompt_hashes(),
        "counts": {
            "responses": result.stats.n_responses,
            "segments": result.stats.n_segments,
            "codes": result.stats.codes_final,
            "families": result.stats.families_final,
            "assignments": result.stats.assignments,
            "snapshots": len(result.snapshot_ids),
            "llm_calls": int(result.stats.llm["calls"]),
        },
    }


def artefact_texts(result: FastLoopResult, config: RunConfig | None = None) -> dict[Path, str]:
    """The exact text each committed fixture file holds, keyed by path."""
    settings = config or golden_config()
    return {
        CODEBOOK_FIXTURE: result.codebook_json() + "\n",
        ASSIGNMENTS_FIXTURE: _dump([a.to_json() for a in result.assignments]),
        SNAPSHOTS_FIXTURE: _dump({"snapshot_ids": list(result.snapshot_ids)}),
        FINDINGS_FIXTURE: _dump(expected_findings(result)),
        MANIFEST_FIXTURE: _dump(expected_manifest(result, settings)),
    }


def read_fixture(path: Path) -> str:
    """Read one committed fixture, with a message that names the way out if it is absent."""
    if not path.exists():
        raise AssertionError(
            f"the golden fixture {path.relative_to(REPO_ROOT)} does not exist.\n"
            f"Generate it deliberately with:\n    {REGENERATE_COMMAND}"
        )
    return path.read_text(encoding="utf-8")


# --------------------------------------------------------------------------- #
# Drift reporting — the failure message is the deliverable
# --------------------------------------------------------------------------- #


def _codebook_headline(expected: str, actual: str) -> list[str]:
    """Name the codes that appeared, vanished or changed, before showing any diff."""
    try:
        before = {c["name"]: c for c in json.loads(expected)["codes"]}
        after = {c["name"]: c for c in json.loads(actual)["codes"]}
    except (ValueError, KeyError, TypeError):
        return ["  (the codebook is not parseable as codebook JSON)"]
    added = sorted(after.keys() - before.keys())
    removed = sorted(before.keys() - after.keys())
    changed = sorted(name for name in before.keys() & after.keys() if before[name] != after[name])
    lines = []
    if added:
        lines.append(f"  codes added   ({len(added)}): {', '.join(added)}")
    if removed:
        lines.append(f"  codes removed ({len(removed)}): {', '.join(removed)}")
    if changed:
        lines.append(f"  codes changed ({len(changed)}): {', '.join(changed)}")
    if not lines:
        lines.append("  the code names are unchanged; only ordering or whitespace moved")
    return lines


def _assignments_headline(expected: str, actual: str) -> list[str]:
    """Name the assignment rows that appeared or vanished."""
    try:
        before = {(r["response_id"], r["code"], r["segment"]) for r in json.loads(expected)}
        after = {(r["response_id"], r["code"], r["segment"]) for r in json.loads(actual)}
    except (ValueError, KeyError, TypeError):
        return ["  (the assignments file is not parseable as assignment rows)"]
    added = sorted(after - before)
    removed = sorted(before - after)
    lines = [f"  rows: {len(before)} expected, {len(after)} produced"]
    for label, rows in (("added", added), ("removed", removed)):
        for response_id, code, segment in rows[:5]:
            lines.append(f"  {label:>7}: response {response_id}  {code}  {segment[:60]!r}")
        if len(rows) > 5:
            lines.append(f"  {label:>7}: ... and {len(rows) - 5} more")
    return lines


def _headline(path: Path, expected: str, actual: str) -> list[str]:
    if path == CODEBOOK_FIXTURE:
        return _codebook_headline(expected, actual)
    if path == ASSIGNMENTS_FIXTURE:
        return _assignments_headline(expected, actual)
    return []


def drift(path: Path, expected: str, actual: str) -> str | None:
    """`None` when the artefact still matches; otherwise the whole failure message.

    A pure function rather than an assertion, so the drift-detection test can call it on
    a deliberately mutated codebook and read what a human would have been shown.
    """
    if expected == actual:
        return None
    name = path.name
    diff = list(
        difflib.unified_diff(
            expected.splitlines(),
            actual.splitlines(),
            fromfile=f"expected/{name}",
            tofile=f"produced/{name}",
            lineterm="",
            n=2,
        )
    )
    truncated = diff[:MAX_DIFF_LINES]
    if len(diff) > MAX_DIFF_LINES:
        truncated.append(f"... {len(diff) - MAX_DIFF_LINES} more diff lines suppressed")
    return "\n".join(
        [
            f"GOLDEN REGRESSION — {name} moved.",
            *_headline(path, expected, actual),
            "",
            *truncated,
            "",
            "The offline pipeline no longer reproduces the committed fixture. Something",
            "changed: a prompt, a threshold, a mock persona, a check severity, the",
            "embedding space, or the corpus fixture. Find out which, and decide whether it",
            "was intended.",
            "If it was intended, rewrite the fixture deliberately and commit the diff:",
            f"    {REGENERATE_COMMAND}",
        ]
    )


# --------------------------------------------------------------------------- #
# 1. The pipeline regression
# --------------------------------------------------------------------------- #


def test_offline_run_reproduces_every_committed_artefact() -> None:
    """The fast loop over the synthetic corpus still produces the committed bytes."""
    result, board = run_golden()
    with board:
        produced = artefact_texts(result)

    messages = [
        message
        for path, actual in produced.items()
        if (message := drift(path, read_fixture(path), actual)) is not None
    ]
    if messages:
        pytest.fail("\n\n".join(messages), pytrace=False)


def test_manifest_matches_the_live_configuration() -> None:
    """The thresholds, prompts and embedding space the fixture assumes are still current.

    Separate from the artefact comparison because it fails for a *different* reason: the
    bytes can still match while a threshold has moved to a value the mock clients happen
    not to exercise, and that is a change a reader of the run report must not miss.
    """
    committed = json.loads(read_fixture(MANIFEST_FIXTURE))
    config = golden_config()
    assert committed["rules"] == config.rules.to_json(), (
        "CodingRules changed since the golden fixture was generated. Re-read the check "
        f"layer, then regenerate on purpose:\n    {REGENERATE_COMMAND}"
    )
    assert committed["embedding"] == config.embedding.to_json()
    assert committed["models"] == config.models.to_json()
    live_prompts = _run_prompt_hashes()
    assert committed["prompts"] == live_prompts, (
        "a prompt template's static text changed. Every coding in the fixture was "
        "produced under the committed wording; regenerate on purpose:\n"
        f"    {REGENERATE_COMMAND}"
    )


def test_a_drifted_codebook_fails_with_a_message_that_names_what_moved() -> None:
    """Construct the drift in-test and read the message a human would be shown.

    The fixture on disk is never touched: the mutation happens to the *produced* side,
    which is what a prompt edit or a threshold change would do.
    """
    expected = read_fixture(CODEBOOK_FIXTURE)
    codebook = Codebook.from_json(json.loads(expected))
    invented = Code(
        id="drift",
        name="drift-invented_by_a_prompt_edit",
        description="a code no committed fixture contains",
        parent_id=None,
        evidence=[Evidence(response_id=203, quote="By 2050")],
        created_in_snapshot="snap-drift",
    )
    drifted = Codebook(codes={**codebook.codes, invented.id: invented})

    message = drift(CODEBOOK_FIXTURE, expected, drifted.to_json_str() + "\n")

    assert message is not None, "a codebook with an extra code must not compare equal"
    assert "GOLDEN REGRESSION" in message
    assert "codes added   (1): drift-invented_by_a_prompt_edit" in message
    assert REGENERATE_COMMAND in message
    assert "+" in message, "the message must carry a unified diff, not only a headline"
    # ...and the unmutated codebook still passes through the same function.
    assert drift(CODEBOOK_FIXTURE, expected, expected) is None


def test_a_dropped_assignment_row_is_named_in_the_failure_message() -> None:
    """The assignments headline says which rows moved, not merely that some did."""
    expected = read_fixture(ASSIGNMENTS_FIXTURE)
    rows = json.loads(expected)
    dropped = rows[0]
    message = drift(ASSIGNMENTS_FIXTURE, expected, _dump(rows[1:]))

    assert message is not None
    assert f"rows: {len(rows)} expected, {len(rows) - 1} produced" in message
    assert dropped["code"] in message
    assert REGENERATE_COMMAND in message


# --------------------------------------------------------------------------- #
# 2. Properties that must never regress, fixture or no fixture
# --------------------------------------------------------------------------- #


def test_two_runs_are_byte_identical() -> None:
    """Acceptance criterion 3, asserted on all four artefacts rather than on the codebook."""
    first, first_board = run_golden("determinism-a")
    second, second_board = run_golden("determinism-b")
    with first_board, second_board:
        left = artefact_texts(first, golden_config("determinism-a"))
        right = artefact_texts(second, golden_config("determinism-b"))

    for path in left:
        if path == MANIFEST_FIXTURE:
            continue  # the manifest records the run id, which is the one thing that differs
        assert left[path] == right[path], f"{path.name} differs between two identical runs"
    assert first.report.to_json() == second.report.to_json()


@pytest.mark.parametrize("seed", [1, 7, 99])
def test_a_shuffled_corpus_produces_the_same_final_codebook(seed: int) -> None:
    """The corpus list order does not reach the pipeline at all.

    NOTE what this does and does not prove. `fast_loop._ordered` sorts by
    `(source, id)` before anything is coded, and every fixture response shares
    `source="synthetic"` — so shuffling the *list* is undone before the first coder call.
    This test therefore pins that the sort is present and load-bearing: delete it and
    this fails. It does NOT by itself establish order independence, because it never
    produces a different processing order. `test_a_genuine_reorder_changes_only_provenance`
    below does that.
    """
    reference = read_fixture(CODEBOOK_FIXTURE)
    shuffled = synthetic_corpus()
    random.Random(seed).shuffle(shuffled)
    assert [r.id for r in shuffled] != response_ids(), "the shuffle did nothing"

    result, board = run_golden(f"shuffle-{seed}", shuffled)
    with board:
        produced = artefact_texts(result)

    assert drift(CODEBOOK_FIXTURE, reference, produced[CODEBOOK_FIXTURE]) is None
    assert produced[ASSIGNMENTS_FIXTURE] == read_fixture(ASSIGNMENTS_FIXTURE)
    assert produced[SNAPSHOTS_FIXTURE] == read_fixture(SNAPSHOTS_FIXTURE)


@pytest.mark.parametrize("seed", [1, 7, 99])
def test_a_genuine_reorder_changes_only_provenance(seed: int) -> None:
    """Order independence, tested where it can actually fail.

    The sort key is `(source, id)`, so the only way to obtain a genuinely different
    processing order is to vary `source`. Ids and text are untouched; each response is
    relabelled into its own source bucket, which permutes the order the fast loop sees.

    What must be identical: the code set, every code's name, description and evidence,
    and every assignment. That is the analytic result, and it is order-independent
    because code identity is content-addressed on the name and evidence is ordered by
    content rather than arrival.

    What legitimately differs: `created_in_snapshot`. It records which batch admitted a
    code, so it is provenance about *this run's* traversal, not a property of the
    codebook. A reorder that left it unchanged would be recording a falsehood. This is
    the honest form of ADR-0021's claim; the earlier wording said the codebook JSON was
    byte-identical under a shuffle, which is false under a reorder the sort does not
    erase.
    """
    baseline, board = run_golden("reorder-baseline")
    with board:
        base_codes = json.loads(baseline.codebook_json())["codes"]
        base_rows = [a.to_json() for a in baseline.assignments]

    corpus = synthetic_corpus()
    order = list(range(len(corpus)))
    random.Random(seed).shuffle(order)
    relabelled = [
        replace(response, source=f"s{position:02d}")
        for position, response in zip(order, corpus, strict=True)
    ]
    assert [r.id for r in sorted(relabelled, key=lambda r: (r.source, r.id))] != [
        r.id for r in corpus
    ], "the relabelling did not change the processing order"

    result, board = run_golden(f"reorder-{seed}", relabelled)
    with board:
        produced_codes = json.loads(result.codebook_json())["codes"]
        produced_rows = [a.to_json() for a in result.assignments]

    def substance(codes: list[dict]) -> list[dict]:
        return [{k: v for k, v in c.items() if k != "created_in_snapshot"} for c in codes]

    assert substance(produced_codes) == substance(base_codes), (
        "a genuine reorder changed the codebook's substance, not merely its provenance"
    )
    assert sorted(map(json.dumps, produced_rows)) == sorted(map(json.dumps, base_rows))


def test_the_fast_loop_never_splits_reparents_or_renames() -> None:
    """The operation set exists, and the fast loop uses none of its restructuring half.

    The predecessor's inability to split or re-parent is the documented cause of codebook
    flattening, so `Operation` carries both — and the fast loop is required never to reach
    for them. This asserts it on the run's own recorded decisions and audit events; B2
    asserts the complementary claim by walking the snapshots.
    """
    assert {"split", "reparent", "rename"} <= set(OPERATION_TYPES), (
        "the restructuring operations are gone from the model, so this test would pass "
        "vacuously"
    )
    result, board = run_golden("no-restructure")
    with board:
        events = set(board.audit("no-restructure").event_names())

    assert set(result.stats.actions) <= {"MERGE", "CREATE"}
    for outcome in result.outcomes:
        for decision in outcome.decisions:
            assert decision["action"] in {"MERGE", "CREATE"}, (
                f"response {outcome.response_id} took action {decision['action']!r}; the "
                "fast loop may only merge or create"
            )
    assert not events & {"code_split", "code_reparented", "code_renamed", "operation_applied"}


def test_no_unexpected_error_finding_appears() -> None:
    """Every ERROR the run emits is one `findings.json` names, and every named one occurs.

    ERRORs in the offline run are planted fixture material — three S2 provenance failures
    from `FABRICATED_QUOTE`, two M3 drops from `MOCKFIT_UNNECESSARY`. A sixth would mean
    the check layer had started failing something it used to pass.
    """
    committed = json.loads(read_fixture(FINDINGS_FIXTURE))
    result, board = run_golden("errors")
    with board:
        produced = expected_findings(result)

    def keys(payload: dict[str, Any]) -> set[tuple[str, str, str, str, int]]:
        return {
            (e["check_id"], e["scope"], e["subject"], e["marker"], e["count"])
            for e in payload["errors"]
        }

    unexpected = keys(produced) - keys(committed)
    missing = keys(committed) - keys(produced)
    assert not unexpected, (
        f"ERROR findings the fixture does not expect: {sorted(unexpected)}.\n"
        "Something in the check layer started failing. If the new ERROR is correct, "
        f"regenerate on purpose:\n    {REGENERATE_COMMAND}"
    )
    assert not missing, (
        f"ERROR findings the fixture expects but the run did not produce: {sorted(missing)}.\n"
        "A check stopped firing, or the planted fixture material stopped tripping it."
    )
    assert produced["summary"] == committed["summary"]
    assert produced["severities"] == committed["severities"]


# --------------------------------------------------------------------------- #
# 3. Concurrent validation against the golden coding
# --------------------------------------------------------------------------- #


def test_the_golden_coding_loads_in_the_row_oriented_shape() -> None:
    """The committed stand-in is exactly what `read_coded_xlsx` would hand back.

    This is the drop-in claim: the PI's workbook and this JSON file produce the same
    `list[Assignment]`, so pointing the harness at the real file changes no code.
    """
    loaded = load_human_coding(HUMAN_CODING_FIXTURE)

    assert loaded == human_assignments(), (
        "the golden human coding no longer agrees with tests.fixtures.assignments. That "
        "fixture is orchestrator-owned; edit tests/fixtures/golden/human_coding.json by "
        "hand to match it — a regeneration must never overwrite a human coding."
    )
    rows = json.loads(HUMAN_CODING_FIXTURE.read_text(encoding="utf-8"))
    assert isinstance(rows, list)
    assert all(sorted(row) == ["code", "response_id", "segment"] for row in rows), (
        "the golden coding must hold exactly the row-oriented assignments shape "
        '[{"response_id": int, "segment": str, "code": str}]'
    )
    assert all(isinstance(row["response_id"], int) for row in rows)


def test_a_workbook_and_the_json_stand_in_load_identically(tmp_path: Path) -> None:
    """The suffix dispatch is real: the same coding, written as xlsx, round-trips."""
    coding = load_human_coding(HUMAN_CODING_FIXTURE)
    workbook = write_coded_xlsx(
        tmp_path / "IndiaHandCoding.xlsx",
        [(a.response_id, a.segment, a.code) for a in coding],
    )
    assert load_human_coding(workbook) == coding


def test_concurrent_validation_reports_the_hand_checked_numbers() -> None:
    """The arithmetic in the module docstring, asserted.

    Hand-derived from the 200-range fixture: the union of the two codings covers
    responses 203, 207, 212, 226, 239, 244, 245, 258 and six code columns, so
    8 x 6 = 48 cells; TP 6, FP 2, FN 2, TN 38; precision = recall = F1 = 0.75;
    observed 44/48, expected 1664/2304, kappa exactly 0.7. The blind spot
    `negative_impacts-lack_of_inclusion` is the human's only unmatched code and it
    falls on responses 207 and 239.
    """
    report = validate_against_golden(
        machine_assignments(),
        human_path=HUMAN_CODING_FIXTURE,
        corpus=corpus_by_id(),
    )
    code = report.code_level

    assert (code.n_responses, code.n_codes, code.n_cells) == (8, 6, 48)
    assert (code.tp, code.fp, code.fn, code.tn) == (6, 2, 2, 38)
    assert code.precision == pytest.approx(0.75)
    assert code.recall == pytest.approx(0.75)
    assert code.f1 == pytest.approx(0.75)
    assert code.observed_agreement == pytest.approx(44 / 48)
    assert code.expected_agreement == pytest.approx(1664 / 2304)
    assert code.kappa == pytest.approx(0.7)

    # Four names shared verbatim match at cosine 1.0; on those columns the machine is
    # perfect, which is what makes the union-vocabulary number above the honest one.
    assert report.matched_only.n_codes == 4
    assert report.matched_only.f1 == pytest.approx(1.0)
    assert report.matched_only.kappa == pytest.approx(1.0)

    # The two lists the write-up quotes.
    assert tuple(u.code for u in report.over_coding) == MACHINE_ONLY
    assert tuple(u.code for u in report.blind_spots) == HUMAN_ONLY
    assert sum(u.n_assignments for u in report.over_coding) == 2
    assert report.blind_spots[0].response_ids == (207, 239)


def test_segment_level_catches_the_right_code_in_the_wrong_place() -> None:
    """Code level scores response 203's job_destruction cell as a hit; segment level must not.

    The human put `negative_impacts-job_destruction` on "Clerical posts in the district
    office vanish quietly" (characters [93, 145) of the normalised response) and the
    machine put it on "no retraining scheme arrives" ([151, 179)). The spans are
    disjoint, so that row scores 0.0 on both sides: five of eight rows match exactly,
    giving weighted precision = recall = 5/8 = 0.625.
    """
    report = validate_against_golden(
        machine_assignments(),
        human_path=HUMAN_CODING_FIXTURE,
        corpus=corpus_by_id(),
    )
    segment = report.segment_level

    assert segment.mode == "span", "the corpus was passed, so spans must be real intervals"
    assert (segment.n_located_human, segment.n_located_machine) == (8, 8)
    assert segment.right_code_wrong_place == 1
    assert segment.weighted_f1 == pytest.approx(0.625)
    assert segment.weighted_f1 < report.code_level.f1, (
        "a coding that puts the right code on the wrong span must score worse at segment "
        "level than at code level, or the two levels are measuring the same thing"
    )


def test_the_harness_validates_the_pipeline_s_own_output() -> None:
    """The composition the real run will use: golden coding versus fast-loop assignments.

    Numbers are not asserted here — they are a property of the mock personas, and job (b)
    already pins those. What is asserted is that the two halves of the harness compose,
    and that the blind-spot list does its job: the loop never invents
    `negative_impacts-lack_of_inclusion`, and a code the machine failed to invent is
    detectable only by this comparison (CODING_RULES, "the one category with no check").
    """
    result, board = run_golden("validation")
    with board:
        report = validate_against_golden(
            result.assignments,
            human_path=HUMAN_CODING_FIXTURE,
            corpus=corpus_by_id(),
        )

    assert report.matching.space_id == "lexical-v1-512"
    assert HUMAN_ONLY[0] in {u.code for u in report.blind_spots}
    assert report.code_level.n_cells > 0
    assert "Concurrent validation" in report.to_markdown()
    assert json.loads(report.to_json_str())["code_level"]["n_cells"] == report.code_level.n_cells


@pytest.mark.skipif(
    HUMAN_CODING_ENV not in os.environ,
    reason=f"{HUMAN_CODING_ENV} is unset: the PI's hand-coding lives outside the repository",
)
def test_real_golden_set_loads_when_supplied() -> None:
    """The day the PI's file lands, this is the test that runs against it.

    Structural only. The numbers it would produce are the study's findings, not a fixture,
    and they belong in the run report rather than in an assertion.
    """
    coding = load_human_coding()
    assert coding, f"{golden_human_coding_path()} contains no assignment rows"
    assert all(isinstance(a, Assignment) for a in coding)
    assert all(a.segment.strip() and a.code.strip() for a in coding)
    assert len({a.code for a in coding}) > 1, "a golden set with one code cannot be validated"


# --------------------------------------------------------------------------- #
# 4. Real-data safety — the constraint that outranks every other test here
# --------------------------------------------------------------------------- #


def _tracked_files() -> list[Path]:
    """Every file the repository plausibly tracks, from the filesystem, no git involved.

    `data/` is pruned to `data/synthetic/`, which is the only child `.gitignore` admits.
    """
    found: list[Path] = []
    stack = [REPO_ROOT]
    while stack:
        directory = stack.pop()
        for entry in sorted(directory.iterdir()):
            if entry.is_symlink():
                continue
            if entry.is_dir():
                if entry.name in IGNORED_DIRS or entry.name.endswith(".egg-info"):
                    continue
                if entry.parent == REPO_ROOT and entry.name == "data":
                    synthetic = entry / "synthetic"
                    if synthetic.is_dir():
                        stack.append(synthetic)
                    continue
                stack.append(entry)
            elif entry.is_file():
                found.append(entry)
    return found


def test_no_real_survey_data_is_committed() -> None:
    """No file with a real-data extension exists anywhere the repository tracks."""
    offenders = [
        path.relative_to(REPO_ROOT)
        for path in _tracked_files()
        if path.suffix.lower() in REAL_DATA_SUFFIXES
    ]
    assert not offenders, (
        f"files with a real-data extension are inside the repository: {offenders}. "
        "Human survey responses must never enter git history; the real corpus lives "
        "outside the repo and stays there."
    )


#: Where the real files live, relative to the repository. Every workbook **and every
#: CSV** directly under these directories is read in place: corpora in either shape and
#: the principal investigator's coding exports, because respondent text arrives in all
#: of them. The first version of this scan read one corpus only, and a fragment of a
#: *different* corpus reached a commit through a code comment (ADR-0030). The second
#: read workbooks only, and the full Process corpus arrived as a CSV — 180 responses
#: the guard could not see, which is ADR-0030's lesson a third time. Set
#: GAF_REAL_CORPUS to a file or a directory (several, separated by os.pathsep) to
#: override.
REAL_DATA_DIRS: tuple[Path, ...] = (
    REPO_ROOT.parent / "data",
    REPO_ROOT.parent / "codebook",
    REPO_ROOT.parent / "Grounded AI Futures" / "data",
    REPO_ROOT.parent / "Grounded AI Futures" / "codebook",
)
#: Kept for callers that name the original corpus explicitly.
REAL_CORPUS_DEFAULT = REAL_DATA_DIRS[2] / "NarrativeState(IndiaSample1-20).xlsx"

#: File types the scan reads. Anything a respondent's words can be stored in: the two
#: workbook shapes, the two delimited shapes, and the principal investigator's codebook
#: Markdown, which the Makefile points the pipeline at as `FULL_DEFINITIONS_MD`. Reading
#: more of them costs nothing and cannot copy anything.
#:
#: `*.docx` is named in `REAL_DATA_SUFFIXES` and in `.gitignore` and is deliberately
#: **not** here, and this was measured rather than assumed. The one real `.docx` is a
#: mirror of the Markdown codebook. A table reader skips a code-name column because
#: `NON_TEXT_HEADERS` names its header; a Word paragraph has no header to skip on, so
#: reading the mirror puts the investigator's **code names and own definitions** into
#: the needle set. Measured: it would charge 20 tracked files, among them
#: `examples/demo-run/codebook.json` and `analysis/clusters.md`, for containing a code
#: name this project legitimately reuses. Neither a code name nor his definition is a
#: respondent's words — the same ruling `_instrument_shingles` makes about the survey
#: question (ADR-0047 decision 3) — and the independent audit's own scanner excluded
#: this file for this reason. Nothing a respondent said is lost: his verbatim
#: highlights reach the scan through the coding exports and the Markdown blockquotes.
REAL_DATA_GLOBS: tuple[str, ...] = ("*.xlsx", "*.xlsm", "*.csv", "*.tsv", "*.md")

def _exporter() -> Any:
    """`scripts/export_results.py`, loaded as a module so its constants can be read.

    The script is not on the import path and is not meant to be; loading it by
    specification is what `tests/test_export_results.py` already does, and the module
    is cached under the same name so it is executed at most once per session.
    """
    cached = sys.modules.get("gaf_export_results")
    if cached is not None:
        return cached
    spec = importlib.util.spec_from_file_location(
        "gaf_export_results", REPO_ROOT / "scripts" / "export_results.py"
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules["gaf_export_results"] = module
    spec.loader.exec_module(module)
    return module


#: Runs of this many characters or more, shared verbatim with a real response, are
#: treated as respondent text. **Loaded from the results exporter, never copied**: this
#: scan governs what may be committed and the exporter's governs what may be shipped,
#: and for one whole branch they disagreed — 30 here against 20 there. A 25-character
#: run of a respondent's words consequently sat in a tracked ADR, invisible, while the
#: stricter guard beside it would have caught it (R2 C-2 and I-1). Twenty is also where
#: the measurement points: 84 of the 342 highlights in the PI's own coding export
#: normalise to under 30 characters, 50 of them into the 20-to-29 band the old value could
#: not see at all. Measured headroom for 20: outside `docs/CODING_RULES.md` the longest
#: respondent run anywhere in the repository is 29 characters, and it is a stock phrase
#: shared with fourteen different responses.
PROVENANCE_SHINGLE: int = _exporter().GUARD_SHINGLE

#: A cell shorter than the shingle is kept **whole** as a needle rather than discarded.
#: This is the floor below which it is not kept at all. The old scan dropped every short
#: cell before the needle set was built, so a short answer was undetectable at any
#: length — a hole distinct from the shingle's, and the one that made the 20-to-29 band
#: unreachable even after lowering the shingle. Twelve is chosen so that a needle is at
#: least a short clause: below it, a cell that survives the "must contain whitespace"
#: rule is a two-word fragment that any English sentence can reproduce by accident.
PROVENANCE_MIN_CELL = 12

#: How many **different** respondents must share a window before it stops being one
#: respondent's words and becomes the corpus's language. A sentence copied out of one
#: answer cannot appear in a second answer written by somebody else, so a window two
#: people both produced was not copied from either. This is the ground R2 applied by
#: hand to thirty-seven of its seventy-five runs; mechanising it is what lets the scan
#: run at twenty characters over a repository whose subject *is* the corpus's subject,
#: where "the difference between" and the PI's own family name "positive impact" are
#: otherwise charged as respondent text. Two is the strictest value that works, and it
#: was regression-tested against the run this pass removed from ADR-0016: all six of
#: that run's windows are unique to a single response, so the rule would still have
#: caught it. See ADR-0047.
PROVENANCE_MIN_SOURCES = 2

#: The one file permitted to contain respondent text, by the principal investigator's
#: explicit decision: it reproduces his own GPTPrompts.docx negative examples, and the
#: coding rules are not defensible without the quotation attached. See ADR-0024 §3 and
#: the README. Every other tracked file must be clean.
PROVENANCE_EXEMPT = {"docs/CODING_RULES.md"}


def _real_sources() -> list[Path]:
    """Every real table the scan can find, read in place. Never copied into the repo."""
    override = os.environ.get("GAF_REAL_CORPUS")
    roots = [Path(p) for p in override.split(os.pathsep)] if override else list(REAL_DATA_DIRS)
    files: list[Path] = []
    for root in roots:
        if root.is_file():
            files.append(root)
        elif root.is_dir():
            for pattern in REAL_DATA_GLOBS:
                files.extend(p for p in root.glob(pattern) if not p.name.startswith("~$"))
    return sorted(set(files))


#: Column headers that mark a column as labels rather than prose: a coding export's
#: tag, id and document-title columns. A cell in such a column is never respondent text.
NON_TEXT_HEADERS = {"id", "tag", "code", "document", "number", "highlight_id", "response_id"}


def _harvest(rows: Iterator[tuple[Any, ...]], texts: list[str]) -> None:
    """Append every prose cell of one table to `texts`, applying the skipping rules.

    Two kinds of cell are *not* respondent text and are skipped: anything in a column
    whose header is in NON_TEXT_HEADERS (the principal investigator's code names and
    document titles, which this repository legitimately reuses), and any cell with no
    whitespace at all, which is a label rather than a sentence. Over-inclusion
    elsewhere only makes the scan stricter; it never copies anything.

    A cell is measured **after** normalisation and kept whenever it reaches
    `PROVENANCE_MIN_CELL`, not `PROVENANCE_SHINGLE`. Keeping only cells long enough to
    carry a shingle is what made short answers undetectable at any length; a short cell
    is instead compared whole by `_provenance_needles` (R2 I-1).
    """
    first = next(rows, None)
    if first is None:
        return
    header = [str(c).strip().casefold() if c is not None else "" for c in first]
    skip = {i for i, h in enumerate(header) if h in NON_TEXT_HEADERS}
    body = rows if skip else (r for r in (first, *rows))
    for row in body:
        for i, cell in enumerate(row):
            if i in skip or not isinstance(cell, str):
                continue
            if not re.search(r"\s", cell):
                continue
            text = re.sub(r"\s+", " ", cell).casefold().strip()
            if len(text) < PROVENANCE_MIN_CELL:
                continue
            texts.append(text)


def _read_real_source(path: Path) -> list[str]:
    """Every prose cell, row or quoted line of one real file, normalised.

    Dispatches on the suffix, because respondent text arrives in five shapes:

    * ``.xlsx`` / ``.xlsm`` — every sheet, every column. The corpus workbooks put the
      response in column B, the numbered shape puts it in column A.
    * ``.csv`` / ``.tsv`` — the numbered CSV puts the response in "All" *and* again in
      "Trimmed", and a coding export puts it under "content". The PI's full
      200-response corpus is a CSV, and a scan that read workbooks only was blind to
      180 of those responses while still passing. ``utf-8-sig`` because his exports
      carry a byte-order mark, which would otherwise leave the first header
      unrecognised and its column unskipped.
    * ``.md`` — the PI's codebook Markdown, which the Makefile hands the pipeline as
      `FULL_DEFINITIONS_MD`. **Only its ``> `` blockquote lines are read**: those are
      his verbatim highlights of respondents. The headings, the tree diagram and the
      definitions beside them are his own writing, and this project reproduces his
      definitions on purpose (ADR-0032) — putting them in the needle set would charge
      it for doing so.
    """
    import csv

    texts: list[str] = []
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as handle:
            _harvest((tuple(row) for row in csv.reader(handle, delimiter=delimiter)), texts)
        return texts
    if suffix == ".md":
        quoted = [
            (line.lstrip()[2:],)
            for line in path.read_text(encoding="utf-8", errors="replace").splitlines()
            if line.lstrip().startswith("> ")
        ]
        # A header row would be consumed, so give the harvest one it will discard.
        _harvest(iter([("quote",), *quoted]), texts)
        return texts

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            _harvest(sheet.iter_rows(values_only=True), texts)
    finally:
        workbook.close()
    return texts


def _real_responses() -> list[str]:
    """Every respondent text this machine can reach, from every real file in every shape."""
    sources = _real_sources()
    if not sources:
        pytest.skip("no real corpus present under REAL_DATA_DIRS; provenance scan skipped")
    texts: list[str] = []
    for path in sources:
        texts.extend(_read_real_source(path))
    return texts


def _instrument_shingles() -> set[str]:
    """Every window of the survey's own question, in every variant `gaf.config` holds.

    The question is the principal investigator's, not a respondent's, and it travels on
    every `Response` record — so it is inside every answer cell the harvest reads, and
    respondents echo it back besides. Left in the needle set it charges three files
    that exist in order to reproduce it: `gaf/config.py`, which holds it as a frozen
    constant; `results/*/01-inputs.md`, which prints what was asked; and
    `examples/demo-run/corpus.json`, which carries it on every record. None of those can
    be reworded, so the correction belongs here. Read from the module rather than
    repeated, so a new variant is covered the day it is added.
    """
    from gaf.config import QUESTION_VARIANTS

    windows: set[str] = set()
    for question in QUESTION_VARIANTS.values():
        text = re.sub(r"\s+", " ", question).casefold().strip()
        for i in range(max(0, len(text) - PROVENANCE_SHINGLE + 1)):
            windows.add(text[i : i + PROVENANCE_SHINGLE])
    return windows


def _provenance_needles(responses: Iterable[str]) -> dict[int, set[str]]:
    """The needle set, bucketed by run length: `{length: {run, ...}}`.

    A response at or above `PROVENANCE_SHINGLE` contributes every run of exactly that
    length. A response below it contributes **itself, whole**, in its own bucket, so
    that it is looked for at the only length at which it can be found rather than not
    at all. This is the same shape `gaf.report.trail._RunWording` uses for the same
    reason, one level up in words instead of characters.

    The survey question is then subtracted: see `_instrument_shingles`.
    """
    needles: dict[int, set[str]] = {}
    for text in responses:
        size = min(PROVENANCE_SHINGLE, len(text))
        if size < PROVENANCE_MIN_CELL:
            continue
        bucket = needles.setdefault(size, set())
        for i in range(len(text) - size + 1):
            bucket.add(text[i : i + size])
    if PROVENANCE_SHINGLE in needles:
        needles[PROVENANCE_SHINGLE] -= _instrument_shingles()
    return {size: bucket for size, bucket in needles.items() if bucket}


def _source_count(window: str, sources: Sequence[str]) -> int:
    """How many different respondents could have produced `window`.

    Counted over *maximal* texts: the PI's 200-response export carries each answer
    twice, under `All` and again under `Trimmed`, and the trimmed cell is a substring
    of the full one. Counting both would let one respondent's sentence pass as two
    people's language, which is the one way this rule could hide a real copy.
    """
    holders = [text for text in sources if window in text]
    if len(holders) < 2:
        return len(holders)
    return len([a for a in holders if not any(a is not b and a in b for b in holders)])


def _provenance_hit(
    body: str,
    needles: dict[int, set[str]] | None = None,
    sources: Sequence[str] | None = None,
) -> tuple[int, int] | None:
    """`(offset, run length)` of the first candidate window in `body`, or `None`.

    The offset is into the whitespace-collapsed, case-folded body, which is what the
    comparison is made against. The length is the bucket that matched — a lower bound
    on the run genuinely shared, not necessarily the maximal one; establishing the
    maximal run means reading the real text, which belongs in an investigation, not in
    a failure message.

    `sources` are the distinct real responses. Given them, a window `PROVENANCE_MIN_SOURCES`
    different respondents share is not a candidate. Omit them and every matching window
    is a candidate, which is what the unit tests below want.
    """
    if needles is None:
        needles = _provenance_needles(_real_responses())
    haystack = re.sub(r"\s+", " ", body).casefold()
    sizes = sorted(needles, reverse=True)
    for i in range(len(haystack)):
        for size in sizes:
            if i + size > len(haystack) or haystack[i : i + size] not in needles[size]:
                continue
            if sources is None or _source_count(haystack[i : i + size], sources) < (
                PROVENANCE_MIN_SOURCES
            ):
                return i, size
            break
    return None


def _offender_line(relative: str, hit: tuple[int, int]) -> str:
    """One offender, as a location and a length — never as the run itself.

    `scripts/export_results.py` was changed for exactly this reason (R2 M-2): a guard
    that prints what it caught puts respondent words into a terminal, a CI log and a
    scrollback buffer, none of which is gitignored and none of which `make scrub`
    empties. The offset is into the file's whitespace-collapsed, case-folded text, so
    the person holding the real corpus can find it and nobody else learns anything.
    """
    offset, length = hit
    return f"{relative}: offset {offset}, run length {length}"


def test_no_tracked_file_contains_real_respondent_text() -> None:
    """The guard the extension check cannot be: provenance, not file type.

    ADR-0024 records why this exists. Real respondent text reached git history once,
    and the check that was run four times matched **file extensions** — which cannot
    see text pasted into a `.py` or a `.md`. This asserts the property that was
    actually violated: that no tracked file shares a run of `PROVENANCE_SHINGLE`
    characters with any real survey response.

    Note what this does that its neighbours cannot. `test_no_real_survey_data_is_committed`
    tests file type. `test_every_golden_segment_is_locatable_in_the_synthetic_corpus`
    tests consistency between two fixtures — if the fixture corpus itself were copied
    from real responses, every golden segment would still be locatable in it, so that
    test cannot establish provenance. Only this one reads the real file and compares.

    Skips when the real corpus is absent, which is CI and every clean clone. That is
    the point: the real file is outside the repository, so this is a check the
    researcher runs on the machine that has the data, before publishing anything.
    """
    responses = _real_responses()
    sources = sorted(set(responses))
    needles = _provenance_needles(responses)

    offenders: list[str] = []
    for path in _tracked_files():
        relative = path.relative_to(REPO_ROOT).as_posix()
        if relative in PROVENANCE_EXEMPT:
            continue
        try:
            body = path.read_text(encoding="utf-8")
        except (UnicodeDecodeError, OSError):
            continue
        hit = _provenance_hit(body, needles=needles, sources=sources)
        if hit is not None:
            offenders.append(_offender_line(relative, hit))

    assert not offenders, (
        "tracked files contain verbatim runs of real respondent text:\n  "
        + "\n  ".join(offenders)
        + "\n\nThe offset is into the file's whitespace-collapsed, case-folded text; the "
        "run itself is deliberately not printed (see _offender_line). This is the defect "
        "ADR-0024 records. Rewrite the text so it is genuinely invented, or — if it is a "
        "deliberate, PI-authorised quotation — add the path to PROVENANCE_EXEMPT and say "
        "so in the README and ADR-0024."
    )


def test_the_exempt_file_is_the_only_one_and_is_declared() -> None:
    """The exemption is narrow, deliberate, and matches what the README tells a reader."""
    assert {"docs/CODING_RULES.md"} == PROVENANCE_EXEMPT
    readme = (REPO_ROOT / "README.md").read_text(encoding="utf-8")
    assert "docs/CODING_RULES.md" in readme, (
        "the one file permitted to quote respondents must be named in the README; a "
        "reader must not have to discover it from a test"
    )


def test_the_provenance_shingle_is_the_exporters_own_and_cannot_drift() -> None:
    """One number, two guards. Two numbers is how C-2 stayed green for a whole branch.

    `scripts/export_results.py` has always scanned at 20 and this scan used 30, so
    anything a respondent said in the 20-to-29 character band was invisible *to the guard
    that governs what may be committed* while being caught by the one that governs what
    may be exported. A quarter of the PI's own highlights normalise to under 30
    characters. This asserts the constant is loaded, not copied.
    """
    assert PROVENANCE_SHINGLE == 20
    assert PROVENANCE_SHINGLE == _exporter().GUARD_SHINGLE


def test_a_cell_shorter_than_the_shingle_still_becomes_a_needle() -> None:
    """A short response is not a safe response; it was merely an unscannable one.

    The old `_harvest` discarded any cell shorter than the shingle *before* the needle
    set was built, so a short answer could not be detected at any length. Everything
    below is invented.
    """
    rows = iter(
        [
            ("id", "All"),
            (1, "the ferry leaves at six"),  # 23 chars, over the shingle
            (2, "the pier is shut"),  # 16 chars, under it, still a needle
            (3, "no idea"),  # 7 chars, under the floor
            (4, "unanswered"),  # no whitespace: a label, not prose
        ]
    )
    texts: list[str] = []
    _harvest(rows, texts)
    assert "the ferry leaves at six" in texts
    assert "the pier is shut" in texts
    assert "no idea" not in texts
    assert "unanswered" not in texts


def test_the_scan_matches_a_short_needle_whole_and_a_long_one_by_shingle() -> None:
    """Below the shingle a needle is compared entire; at or above it, by run.

    Invented text throughout. The negative control is what stops the scan degenerating
    into "everything matches".
    """
    needles = _provenance_needles(["the pier is shut", "the ferry leaves at six today"])
    assert _provenance_hit("notice: the pier is shut until spring", needles=needles) is not None
    assert _provenance_hit("we heard the ferry leaves at six today", needles=needles) is not None
    # One character short of the whole short needle, and nothing of the long one.
    assert _provenance_hit("the pier is shu", needles=needles) is None
    assert _provenance_hit("a wholly unrelated sentence about bicycles", needles=needles) is None


def test_the_scan_reads_every_file_type_the_project_calls_real_data(tmp_path: Path) -> None:
    """`REAL_DATA_SUFFIXES` and `.gitignore` name five types; the scan read three.

    A `.tsv` and the principal investigator's codebook Markdown — which the Makefile
    points the pipeline at as `FULL_DEFINITIONS_MD` — were outside the scan entirely
    (R2 M-3). Everything below is invented.
    """
    assert {"*.xlsx", "*.xlsm", "*.csv", "*.tsv", "*.md"} <= set(REAL_DATA_GLOBS)

    (tmp_path / "corpus.tsv").write_text(
        "id\tAll\n1\tthe ferry leaves at six every morning\n", encoding="utf-8"
    )
    (tmp_path / "codebook.md").write_text(
        "## harbour-timetable\n\nSailings and their hours.\n\n"
        "> the pier is shut until spring\n",
        encoding="utf-8",
    )
    texts = _read_real_source(tmp_path / "corpus.tsv") + _read_real_source(
        tmp_path / "codebook.md"
    )
    assert "the ferry leaves at six every morning" in texts
    # A blockquote is the PI's verbatim highlight of a respondent; his own definition
    # beside it is his writing, and this project reproduces it on purpose (ADR-0032).
    assert "the pier is shut until spring" in texts
    assert not any("sailings and their hours" in t for t in texts)


def test_the_survey_instrument_is_not_one_respondents_words() -> None:
    """The question is the PI's, not an answer, and it travels on every response record.

    Respondents echo the question back, so every needle set built from their answers
    contains it. Leaving it in makes the guard charge `gaf/config.py` — where the
    question is a constant — and `results/*/01-inputs.md` and `examples/demo-run/corpus.json`,
    which reproduce it because reproducing the instrument is what those files are for.
    None of the three can be reworded, so the correction belongs in the needle set.
    """
    instrument = _instrument_shingles()
    assert instrument, "no question constant was found in gaf.config"
    sample = re.sub(r"\s+", " ", QUESTION_V2).casefold().strip()
    window = sample[:PROVENANCE_SHINGLE]
    assert window in instrument
    # A response that echoes the question contributes no needle the question already is,
    # so a file that reproduces the instrument is not charged for it. (Where a respondent
    # runs his own words up against the question, the window that straddles the join is
    # his and stays a needle: the subtraction is of the question, not of its neighbours.)
    needles = _provenance_needles([sample, f"i agree. {sample} i have no answer."])
    assert window not in needles.get(PROVENANCE_SHINGLE, set())
    assert _provenance_hit(QUESTION_V2, needles=needles) is None


def test_a_window_two_respondents_share_is_the_corpus_language() -> None:
    """Multiplicity, mechanised: a sentence copied from one answer cannot be in two.

    This is the rule R2 applied by hand to thirty-seven of its seventy-five runs, and
    the only one that lets the guard run at twenty characters over a repository whose
    subject is the corpus's own subject. Everything below is invented.
    """
    shared = "the pier is shut for the winter months"
    lone = "the harbour master keeps a ledger of every crossing"
    sources = [f"{shared} again", f"we were told {shared}", lone]
    needles = _provenance_needles(sources)
    # Two different responses carry it, so it is the language, not one person's words.
    assert _provenance_hit(shared, needles=needles, sources=sources) is None
    # One response carries it, so it is a candidate and a person must rule it.
    assert _provenance_hit(lone, needles=needles, sources=sources) is not None
    # Without a source list there is no multiplicity test and every window is a candidate.
    assert _provenance_hit(shared, needles=needles) is not None


def test_a_provenance_offender_line_carries_no_respondent_text() -> None:
    """The failure message is a location and a length, never the run it caught.

    The same rule the exporter's guard follows (R2 M-2): a message that prints the text
    it caught puts respondent words into a terminal, a CI log and a scrollback buffer,
    none of which is gitignored and none of which `make scrub` empties.
    """
    needles = _provenance_needles(["the pier is shut"])
    hit = _provenance_hit("notice: the pier is shut until spring", needles=needles)
    assert hit is not None
    line = _offender_line("docs/EXAMPLE.md", hit)
    assert "pier" not in line and "shut" not in line
    assert "docs/EXAMPLE.md" in line
    assert "offset 8" in line and "run length 16" in line


def test_the_gitignore_guards_are_still_in_place() -> None:
    """The property above holds today; these lines are what keep it holding tomorrow."""
    ignore = (REPO_ROOT / ".gitignore").read_text(encoding="utf-8").splitlines()
    lines = {line.strip() for line in ignore}
    assert "data/*" in lines and "!data/synthetic/" in lines
    for suffix in sorted(REAL_DATA_SUFFIXES):
        assert f"*{suffix}" in lines, f"{suffix} is no longer excluded repository-wide"


def test_every_golden_fixture_is_json() -> None:
    """A strict allowlist on the golden directory: no spreadsheet can land here by accident.

    `.gitignore` un-ignores `tests/fixtures/**/*.csv`, so a CSV dropped into this directory
    *would* be tracked. The allowlist closes that door.
    """
    files = sorted(p for p in GOLDEN_DIR.rglob("*") if p.is_file())
    assert files, "the golden fixture directory is empty"
    for path in files:
        assert path.suffix == ".json", (
            f"{path.relative_to(REPO_ROOT)} is not JSON. Everything in the golden "
            "directory is a synthetic, human-readable artefact."
        )
        json.loads(path.read_text(encoding="utf-8"))


def test_the_golden_corpus_is_the_synthetic_fixture_and_not_the_real_sample() -> None:
    """Response ids in every golden fixture come from the invented 14-response corpus."""
    synthetic_ids = set(response_ids())
    assert len(synthetic_ids) == 14, "the synthetic corpus changed shape"
    assert not synthetic_ids & REAL_SAMPLE_ONLY_IDS

    manifest = json.loads(read_fixture(MANIFEST_FIXTURE))
    assert manifest["response_ids"] == response_ids()

    rows = json.loads(read_fixture(ASSIGNMENTS_FIXTURE))
    assignment_ids = {row["response_id"] for row in rows}
    coding_ids = {a.response_id for a in load_human_coding(HUMAN_CODING_FIXTURE)}
    for label, ids in (("assignments", assignment_ids), ("human coding", coding_ids)):
        assert ids <= synthetic_ids, (
            f"the golden {label} references response ids outside the synthetic corpus: "
            f"{sorted(ids - synthetic_ids)}"
        )
        assert not ids & REAL_SAMPLE_ONLY_IDS


def test_every_golden_segment_is_locatable_in_the_synthetic_corpus() -> None:
    """The proof that no fixture derives from a respondent's words.

    Every quoted segment — in the expected assignments and in the human golden coding —
    must occur verbatim in `normalise(response.content)` of the invented corpus. A segment
    copied or paraphrased from the real survey could not satisfy this.
    """
    corpus = {rid: normalise(r.content) for rid, r in corpus_by_id().items()}
    rows = [Assignment.from_json(row) for row in json.loads(read_fixture(ASSIGNMENTS_FIXTURE))]
    rows.extend(load_human_coding(HUMAN_CODING_FIXTURE))

    for assignment in rows:
        haystack = corpus.get(assignment.response_id)
        assert haystack is not None, f"response {assignment.response_id} is not in the corpus"
        assert normalise(assignment.segment) in haystack, (
            f"segment {assignment.segment!r} on response {assignment.response_id} does not "
            "occur in the synthetic corpus, so it came from somewhere else"
        )


# --------------------------------------------------------------------------- #
# Deliberate regeneration
# --------------------------------------------------------------------------- #


def regenerate() -> list[Path]:
    """Rewrite the expected artefacts from a fresh run. Never touches the human coding."""
    result, board = run_golden()
    with board:
        texts = artefact_texts(result)
    GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
    for path, text in texts.items():
        path.write_text(text, encoding="utf-8")
    return sorted(texts)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m tests.test_golden",
        description=(
            "Rewrite the golden pipeline-regression fixtures from a fresh offline run. "
            "Do this only when a change to the pipeline was intended, and commit the "
            "resulting diff as part of that change."
        ),
    )
    parser.add_argument(
        "--regenerate",
        action="store_true",
        required=True,
        help="required; without it this command does nothing, so the fixture cannot be "
        "rewritten by accident",
    )
    parser.parse_args(argv)
    for path in regenerate():
        print(f"wrote {path.relative_to(REPO_ROOT)}")
    print("\nReview the diff before committing: these bytes are the regression baseline.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
