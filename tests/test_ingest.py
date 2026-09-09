"""Ingest: spreadsheet readers, encoding repair, and canonical corpus assembly.

The acceptance criterion this file guards is the one the whole pipeline rests on:
**every record carries an id, a question and content**, with the survey question --
which appears nowhere in the source file -- attached to all of them.

Every workbook here is built in-test. No human survey response appears in this
repository; the real file's *shape* is reproduced faithfully (non-contiguous ids, a
used range wider than the labelled columns, encoding damage) without reproducing any
of its content.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from gaf.config import QUESTION_V2, QUESTION_V3, RunConfig
from gaf.ids import content_hash
from gaf.ingest.corpus import (
    META_ENCODING_REPAIRED,
    IngestReport,
    attach_question,
    corpus_content_hash,
    load_corpus,
    load_corpus_with_report,
    read_corpus_json,
    response_content_hash,
    source_from_path,
    write_corpus_json,
)
from gaf.ingest.xlsx import (
    SpreadsheetFormatError,
    read_coded_xlsx,
    read_narrative_state_xlsx,
    repair_mojibake,
)
from gaf.models import Assignment, Response
from gaf.textnorm import normalise
from tests.fixtures.xlsx import (
    CODED_HEADER_VARIANTS,
    SAMPLE_ROWS,
    write_coded_xlsx,
    write_narrative_state_xlsx,
)

# The real file's ids: they start at 9 and skip. Anything assuming range(len(rows))
# must break here rather than on the PI's data.
NON_CONTIGUOUS_IDS = (9, 10, 12, 16, 17, 18, 21, 25, 26, 27, 31, 39, 41, 49, 50, 53, 54, 56, 57, 58)

# UTF-8 curly quotes decoded once as MacRoman. U+201C is E2 80 9C, which MacRoman
# renders as U+201A U+00C4 U+00FA; U+201D is E2 80 9D -> U+201A U+00C4 U+00F9.
MOJIBAKE_OPEN = "\u201a\u00c4\u00fa"
MOJIBAKE_CLOSE = "\u201a\u00c4\u00f9"
MOJIBAKE_TEXT = f"He said {MOJIBAKE_OPEN}Not now. I am tired{MOJIBAKE_CLOSE} and walked away."
REPAIRED_TEXT = "He said “Not now. I am tired” and walked away."


def _rows(count: int) -> list[tuple[int, str]]:
    """`count` rows with the real file's id sequence."""
    return [(rid, f"Response text number {rid} about AI in 2050.") for rid in NON_CONTIGUOUS_IDS[:count]]


# --------------------------------------------------------------------------- #
# Raw-response reader: headers, padding, blank rows, ids
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("header", [("Number", "Response"), ("No", "Response")])
def test_both_id_header_spellings_parse(tmp_path: Path, header: tuple[str, str]) -> None:
    """The PI's prompt versions disagree on "No" vs "Number"; both must read."""
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx", header=header)
    rows = read_narrative_state_xlsx(path)
    assert [r.response_id for r in rows] == [rid for rid, _ in SAMPLE_ROWS]


def test_header_case_and_whitespace_are_tolerated(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx", header=("  NUMBER ", "response  "))
    assert len(read_narrative_state_xlsx(path)) == len(SAMPLE_ROWS)


def test_missing_required_column_raises_naming_file_and_headers(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx", header=("Number", "Answer"))
    with pytest.raises(SpreadsheetFormatError) as excinfo:
        read_narrative_state_xlsx(path)
    message = str(excinfo.value)
    assert "raw.xlsx" in message  # names the file
    assert "Answer" in message  # names the headers actually found
    assert "response" in message  # names what it needed


def test_blank_trailing_rows_are_ignored(tmp_path: Path) -> None:
    """openpyxl routinely yields trailing all-empty rows; they are not records."""
    path = tmp_path / "trailing.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Number", "Response"])
    for number, text in _rows(3):
        sheet.append([number, text])
    for _ in range(4):
        sheet.append([None, None])
    sheet.append(["", "   "])  # blank-but-present strings count as blank too
    workbook.save(path)

    rows = read_narrative_state_xlsx(path)
    assert [r.response_id for r in rows] == list(NON_CONTIGUOUS_IDS[:3])


def test_trailing_unlabelled_empty_columns_are_dropped(tmp_path: Path) -> None:
    """The real file's used range is A1:D21: every row arrives as a 4-tuple.

    Positional unpacking dies on this. Column resolution by header name must not.
    """
    path = tmp_path / "padded.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Number", "Response", None, None])
    for number, text in _rows(20):
        sheet.append([number, text, None, None])
    workbook.save(path)

    rows = read_narrative_state_xlsx(path)
    assert sheet.max_column == 4  # the padding really is there
    assert len(rows) == 20
    assert [r.response_id for r in rows] == list(NON_CONTIGUOUS_IDS)
    assert all(r.content for r in rows)


def test_unlabelled_column_carrying_data_raises_rather_than_guessing(tmp_path: Path) -> None:
    path = tmp_path / "mystery.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Number", "Response", None])
    sheet.append([9, "text", "something nobody labelled"])
    workbook.save(path)

    with pytest.raises(SpreadsheetFormatError, match="no header but carries data"):
        read_narrative_state_xlsx(path)


def test_non_contiguous_ids_survive_ingest(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx", rows=_rows(20))
    rows = read_narrative_state_xlsx(path)
    assert [r.response_id for r in rows] == list(NON_CONTIGUOUS_IDS)
    assert rows[0].response_id == 9  # not 0, not 1


def test_row_with_id_but_no_text_raises(tmp_path: Path) -> None:
    path = tmp_path / "hole.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Number", "Response"])
    sheet.append([9, "fine"])
    sheet.append([10, None])
    workbook.save(path)
    with pytest.raises(SpreadsheetFormatError, match="empty Response cell"):
        read_narrative_state_xlsx(path)


def test_duplicate_ids_raise(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(
        tmp_path / "dup.xlsx", rows=[(9, "one"), (10, "two"), (9, "again")]
    )
    with pytest.raises(SpreadsheetFormatError, match="already appeared"):
        read_narrative_state_xlsx(path)


def test_missing_file_raises_a_named_error(tmp_path: Path) -> None:
    with pytest.raises(SpreadsheetFormatError, match="no such spreadsheet"):
        read_narrative_state_xlsx(tmp_path / "absent.xlsx")


# --------------------------------------------------------------------------- #
# Coded workbooks (the golden set)
# --------------------------------------------------------------------------- #

CODED_ROWS = (
    (9, "AI will help people in many ways", "positive_impacts-problem-solving"),
    (9, "reduce the effort needed for daily work", "positive_impacts-efficiency"),
    (12, "it will keep improving", "future-inevitability"),
)


@pytest.mark.parametrize("header", CODED_HEADER_VARIANTS)
def test_coded_workbook_parses_under_both_header_variants(
    tmp_path: Path, header: tuple[str, str, str]
) -> None:
    path = write_coded_xlsx(tmp_path / "coded.xlsx", rows=CODED_ROWS, header=header)
    assignments = read_coded_xlsx(path)
    assert assignments == [Assignment(rid, segment, code) for rid, segment, code in CODED_ROWS]


def test_coded_workbook_segment_is_the_segment_not_the_response(tmp_path: Path) -> None:
    """The "Response" column of a coded workbook holds the coded span, and repeats ids."""
    path = write_coded_xlsx(tmp_path / "coded.xlsx", rows=CODED_ROWS)
    assignments = read_coded_xlsx(path)
    assert [a.response_id for a in assignments] == [9, 9, 12]
    assert assignments[0].segment != assignments[1].segment


def test_coded_workbook_missing_code_column_raises(tmp_path: Path) -> None:
    path = write_coded_xlsx(
        tmp_path / "coded.xlsx", rows=CODED_ROWS, header=("No", "Response", "Theme")
    )
    with pytest.raises(SpreadsheetFormatError, match="required column missing"):
        read_coded_xlsx(path)


def test_coded_workbook_tolerates_padding_and_blank_rows(tmp_path: Path) -> None:
    path = tmp_path / "coded_padded.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["No", "Response", "Code", None])
    for rid, segment, code in CODED_ROWS:
        sheet.append([rid, segment, code, None])
    sheet.append([None, None, None, None])
    workbook.save(path)
    assert len(read_coded_xlsx(path)) == 3


def test_coded_workbook_incomplete_row_raises(tmp_path: Path) -> None:
    path = write_coded_xlsx(
        tmp_path / "coded.xlsx", rows=[(9, "a segment", "a-code"), (10, "another segment", "")]
    )
    with pytest.raises(SpreadsheetFormatError, match="incomplete assignment"):
        read_coded_xlsx(path)


# --------------------------------------------------------------------------- #
# Encoding repair
# --------------------------------------------------------------------------- #


def test_repair_undoes_utf8_read_as_a_legacy_codepage() -> None:
    assert repair_mojibake(MOJIBAKE_TEXT) == REPAIRED_TEXT


def test_repair_leaves_a_genuine_ellipsis_alone() -> None:
    """U+2026 is in the suspect set, but no re-decoding improves it, so it stays."""
    text = "The list goes on… and on."
    assert repair_mojibake(text) == text


def test_repair_leaves_a_lone_low_quote_alone() -> None:
    """A single U+201A is suspicious but not damage; a repair must strictly improve."""
    text = "A stray mark \u201a in otherwise clean text."
    assert repair_mojibake(text) == text


def test_repair_leaves_clean_text_untouched() -> None:
    text = "Plain ASCII text about AI in 2050."
    assert repair_mojibake(text) is text


def test_repair_is_idempotent() -> None:
    once = repair_mojibake(MOJIBAKE_TEXT)
    assert repair_mojibake(once) == once


def test_repair_happens_at_ingest_and_is_reported(tmp_path: Path) -> None:
    """A silent correction of the PI's data is exactly what transparency forbids."""
    path = write_narrative_state_xlsx(
        tmp_path / "NarrativeState(IndiaSample1-20).xlsx",
        rows=[(9, "Clean text."), (57, MOJIBAKE_TEXT)],
    )
    records, report = load_corpus_with_report(path)
    by_id = {r.id: r for r in records}

    assert by_id[57].content == REPAIRED_TEXT
    assert by_id[57].meta[META_ENCODING_REPAIRED] is True
    assert by_id[9].meta == {}  # untouched records carry no flag
    assert report.repaired_response_ids == (57,)
    assert report.repaired_count == 1
    assert "1 of 2 required encoding repair" in report.summary()


def test_repaired_text_normalises_to_ascii(tmp_path: Path) -> None:
    """Repair then normalise leaves nothing exotic for a span to index into."""
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx", rows=[(57, MOJIBAKE_TEXT)])
    (record,) = load_corpus(path)
    assert all(ord(ch) < 128 for ch in normalise(record.content))


# --------------------------------------------------------------------------- #
# Corpus assembly: the question, the source, the hashes
# --------------------------------------------------------------------------- #


def test_three_row_workbook_yields_exactly_three_records(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx")  # SAMPLE_ROWS: ids 9, 10, 12
    records = load_corpus(path)
    assert len(records) == 3
    assert [r.id for r in records] == [9, 10, 12]


def test_every_record_carries_id_question_and_content(tmp_path: Path) -> None:
    """The acceptance criterion, stated as an assertion."""
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx", rows=_rows(20))
    records = load_corpus(path)
    assert len(records) == 20
    for record in records:
        assert isinstance(record.id, int)
        assert record.question == QUESTION_V2
        assert record.content
    assert [r.id for r in records] == list(NON_CONTIGUOUS_IDS)


def test_question_variant_selects_the_question(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx")
    default = load_corpus(path)
    v3 = load_corpus(path, config=RunConfig().with_(question_variant="v3"))

    assert {r.question for r in default} == {QUESTION_V2}
    assert {r.question for r in v3} == {QUESTION_V3}
    assert QUESTION_V2 != QUESTION_V3


def test_explicit_question_overrides_the_variant(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx")
    records, report = load_corpus_with_report(path, question="A bespoke question?")
    assert {r.question for r in records} == {"A bespoke question?"}
    assert report.question_variant == "custom"


def test_source_is_derived_from_the_file_name() -> None:
    assert source_from_path("NarrativeState(IndiaSample1-20).xlsx") == "india_sample_1_20"
    assert source_from_path(Path("/tmp/NarrativeState(IndiaSample21-40).xlsx")) == "india_sample_21_40"
    assert source_from_path("plain_corpus.json") == "plain_corpus"


def test_source_travels_onto_every_record(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "NarrativeState(IndiaSample1-20).xlsx")
    records, report = load_corpus_with_report(path)
    assert report.source == "india_sample_1_20"
    assert {r.source for r in records} == {"india_sample_1_20"}


def test_explicit_source_wins(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx")
    records = load_corpus(path, source="pilot_wave")
    assert {r.source for r in records} == {"pilot_wave"}


def test_content_hash_is_content_addressed(tmp_path: Path) -> None:
    path = write_narrative_state_xlsx(tmp_path / "raw.xlsx")
    records = load_corpus(path)
    assert response_content_hash(records[0]) == content_hash(records[0].content)
    assert response_content_hash(records[0]) != response_content_hash(records[1])
    # The corpus hash is a pure function of the corpus, not of the load.
    assert corpus_content_hash(records) == corpus_content_hash(load_corpus(path))


def test_corpus_order_is_deterministic_and_independent_of_row_order(tmp_path: Path) -> None:
    forward = write_narrative_state_xlsx(
        tmp_path / "a.xlsx", rows=[(9, "nine"), (12, "twelve"), (41, "forty one")]
    )
    shuffled = write_narrative_state_xlsx(
        tmp_path / "b.xlsx", rows=[(41, "forty one"), (9, "nine"), (12, "twelve")]
    )
    first = load_corpus(forward, source="s")
    second = load_corpus(shuffled, source="s")
    assert [r.id for r in first] == [9, 12, 41]
    assert first == second
    assert corpus_content_hash(first) == corpus_content_hash(second)


def test_unsupported_suffix_raises(tmp_path: Path) -> None:
    path = tmp_path / "corpus.txt"
    path.write_text("nope", encoding="utf-8")
    with pytest.raises(ValueError, match="unsupported corpus format"):
        load_corpus(path)


def test_attach_question_replaces_it_on_every_record() -> None:
    records = [Response(id=9, question="", content="text", source="s")]
    assert attach_question(records, QUESTION_V3)[0].question == QUESTION_V3


# --------------------------------------------------------------------------- #
# JSON round-trip
# --------------------------------------------------------------------------- #


def test_corpus_json_round_trip_is_lossless_and_deterministic(tmp_path: Path) -> None:
    xlsx = write_narrative_state_xlsx(
        tmp_path / "raw.xlsx", rows=[(9, "Clean text."), (57, MOJIBAKE_TEXT)]
    )
    records = load_corpus(xlsx)

    out = write_corpus_json(records, tmp_path / "corpus.json")
    first_bytes = out.read_bytes()
    write_corpus_json(records, out)
    assert out.read_bytes() == first_bytes  # two writes, identical bytes

    reloaded = read_corpus_json(out)
    assert reloaded == records
    assert reloaded[1].meta[META_ENCODING_REPAIRED] is True  # provenance survives


def test_corpus_json_loads_through_load_corpus(tmp_path: Path) -> None:
    xlsx = write_narrative_state_xlsx(tmp_path / "raw.xlsx")
    records = load_corpus(xlsx)
    out = write_corpus_json(records, tmp_path / "corpus.json")
    assert load_corpus(out) == records


def test_corpus_json_without_a_question_gets_the_configured_one(tmp_path: Path) -> None:
    path = tmp_path / "bare.json"
    path.write_text(
        json.dumps([{"id": 9, "question": "", "content": "text", "source": "s"}]),
        encoding="utf-8",
    )
    (record,) = load_corpus(path, config=RunConfig().with_(question_variant="v3"))
    assert record.question == QUESTION_V3


def test_ingest_report_serialises() -> None:
    report = IngestReport(
        path="x.xlsx", source="s", question_variant="v2", record_count=20,
        repaired_response_ids=(57,),
    )
    assert report.to_json()["repaired_response_ids"] == [57]
    assert json.dumps(report.to_json())  # JSON-serialisable, for the run report


# --------------------------------------------------------------------------- #
# The two shapes the PI's Process sample arrived in (2026-09-09). Synthetic text only.
# --------------------------------------------------------------------------- #

from gaf.ingest.corpus import (  # noqa: E402 - grouped with the tests that use them
    META_DUPLICATE_OF_NUMBER,
    META_DUPLICATE_ORDINAL,
    META_SOURCE_NUMBER,
)
from gaf.ingest.xlsx import (  # noqa: E402
    DUPLICATE_ID_STRIDE,
    read_highlights_xlsx,
    read_numbered_column_xlsx,
    sniff_numbered_column,
)
from tests.fixtures.corpus import synthetic_corpus  # noqa: E402


def _numbered_workbook(path: Path, cells: list[str]) -> Path:
    """One column, no header, each cell 'NN. text' — the PI's corpus-sample shape."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IndiaProcess"
    for cell in cells:
        sheet.append([cell])
    workbook.save(path)
    return path


def test_numbered_column_is_detected_and_the_headed_shape_is_not(tmp_path: Path) -> None:
    numbered = _numbered_workbook(tmp_path / "n.xlsx", ["5. Clinics get support.", "7. Posts vanish."])
    headed = write_narrative_state_xlsx(tmp_path / "h.xlsx")
    assert sniff_numbered_column(numbered) is True
    assert sniff_numbered_column(headed) is False


def test_numbered_column_splits_the_number_off_and_keeps_file_order(tmp_path: Path) -> None:
    path = _numbered_workbook(
        tmp_path / "n.xlsx",
        ["11. Clinics get diagnostic support.", "14) Posts in the office vanish.", " 19.   By then the reservoirs run themselves."],
    )
    rows = read_numbered_column_xlsx(path)
    assert [r.response_id for r in rows] == [11, 14, 19]
    assert [r.source_number for r in rows] == [11, 14, 19]
    assert all(r.duplicate_ordinal == 0 for r in rows)
    assert rows[0].content == "Clinics get diagnostic support."
    assert rows[2].content == "By then the reservoirs run themselves."  # prefix and padding stripped


def test_a_repeated_number_keeps_both_responses_and_records_the_fact(tmp_path: Path) -> None:
    """The real file numbers two different responses 44. Neither may be dropped or merged.

    The headed reader raises on a duplicate id because two rows claiming one respondent
    cannot both be that respondent. Here the number is the PI's own label and both
    texts are genuine, so the second is offset to stay unique and the original number
    is kept on the row — reported, never silently repaired.
    """
    path = _numbered_workbook(
        tmp_path / "dup.xlsx",
        ["44. The first response numbered forty-four.", "44. A different response, also forty-four.", "45. The next one."],
    )
    rows = read_numbered_column_xlsx(path)
    assert [r.response_id for r in rows] == [44, 44 + DUPLICATE_ID_STRIDE, 45]
    assert [r.source_number for r in rows] == [44, 44, 45]
    assert [r.duplicate_ordinal for r in rows] == [0, 1, 0]
    assert rows[1].content.startswith("A different response")


def test_load_corpus_autodetects_the_numbered_shape_and_reports_duplicates(tmp_path: Path) -> None:
    path = _numbered_workbook(
        tmp_path / "corpus.xlsx",
        ["3. Rural clinics get support.", "3. A second three.", "8. Posts vanish quietly."],
    )
    responses, report = load_corpus_with_report(path)
    assert [r.id for r in responses] == [3, 8, 3 + DUPLICATE_ID_STRIDE]  # (source, id) order
    assert all(r.question for r in responses), "the survey question is attached to every record"
    dup = next(r for r in responses if r.id == 3 + DUPLICATE_ID_STRIDE)
    assert dup.meta[META_SOURCE_NUMBER] == 3
    assert dup.meta[META_DUPLICATE_OF_NUMBER] == 3
    assert dup.meta[META_DUPLICATE_ORDINAL] == 1
    assert report.duplicate_numbers == ((3, 3 + DUPLICATE_ID_STRIDE),)
    assert "number 3 appears again" in report.summary()


def test_a_two_column_sheet_without_headers_is_still_refused(tmp_path: Path) -> None:
    """Autodetection must not turn the reader into a guesser."""
    path = tmp_path / "two.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["11. text", "extra"])
    workbook.save(path)
    with pytest.raises(SpreadsheetFormatError):
        load_corpus_with_report(path)


def _highlights_workbook(path: Path, rows: list[tuple[str, str, str, str]]) -> Path:
    """A coding-tool export: id | document | tag | content. No response number."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "highlights"
    sheet.append(["id", "document", "tag", "content"])
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def test_highlights_are_placed_by_locating_their_text(tmp_path: Path) -> None:
    """The export names no response; every highlight is placed by where its text occurs.

    Built on the synthetic corpus: exact substrings of three different responses, the
    same segment coded twice (the PI's two-codes-per-segment rule in the data), one
    single word that appears in two responses (ambiguous, excluded), and one string
    that appears nowhere (unlocated, excluded).
    """
    corpus = synthetic_corpus()
    by_id = {r.id: r for r in corpus}
    seg203 = "Rural clinics get diagnostic support"
    seg207 = "Firms will hand ticket triage and first-draft paperwork to software"
    seg253 = "nobody in the depot will remember how the rota used to be built"
    assert seg203 in by_id[203].content and seg207 in by_id[207].content and seg253 in by_id[253].content
    shared_word = "will"  # present in many responses -> ambiguous
    assert sum(1 for r in corpus if shared_word in r.content.split()) > 1

    path = _highlights_workbook(
        tmp_path / "hl.xlsx",
        [
            ("1", "Doc (1-100)", "positive_impacts-healthcare", seg203),
            ("2", "Doc (1-100)", "negative_impacts-job_destruction", seg207),
            ("3", "Doc (1-100)", "negative_impacts-dependence", seg253),
            ("3", "Doc (1-100)", "future-uncertainty", seg253),  # same segment, second code
            ("4", "Doc (1-100)", "future-inevitability", shared_word),
            ("5", "Doc (1-100)", "future-superintelligence", "a sentence that occurs in no response at all"),
        ],
    )
    assignments, report = read_highlights_xlsx(path, corpus)

    assert report.total == 6
    assert (report.exact, report.fuzzy, report.ambiguous, report.unlocated) == (4, 0, 1, 1)
    assert report.mapped == 4
    assert [a.response_id for a in assignments] == [203, 207, 253, 253]
    assert [a.code for a in assignments][2:] == ["negative_impacts-dependence", "future-uncertainty"]
    assert report.per_response == {203: 1, 207: 1, 253: 2}
    outcomes = {m.highlight_id: m.outcome for m in report.mappings}
    assert outcomes["4"] == "ambiguous" and len(next(m for m in report.mappings if m.highlight_id == "4").candidates) > 1
    assert outcomes["5"] == "unlocated"
    assert "4 mapped" in report.summary() and "1 ambiguous" in report.summary()


def test_highlights_accept_code_and_text_as_header_synonyms(tmp_path: Path) -> None:
    corpus = synthetic_corpus()
    path = tmp_path / "syn.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Code", "Text"])
    sheet.append(["positive_impacts-healthcare", "Rural clinics get diagnostic support"])
    workbook.save(path)
    assignments, report = read_highlights_xlsx(path, corpus)
    assert len(assignments) == 1 and assignments[0].response_id == 203
    assert report.mappings[0].highlight_id == "0"  # no id column -> row offset


def test_highlights_missing_tag_column_raises_naming_the_file(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "content"])
    sheet.append(["1", "anything"])
    workbook.save(path)
    with pytest.raises(SpreadsheetFormatError, match=r"bad\.xlsx"):
        read_highlights_xlsx(path, synthetic_corpus())
