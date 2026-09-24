"""Wave 0 gate: the frozen contracts behave as specified, and the fixtures are sound.

This file is the smoke test the brief requires to be green before any Wave 1 agent
starts. It asserts three kinds of thing:

1. the contracts do what `docs/ARCHITECTURE.md` and the brief say they do;
2. the invariants that later waves depend on hold (deterministic serialisation, the
   append-only audit log, non-destructive fail-safes);
3. the fixtures actually contain the violations they claim to, so a later agent that
   writes a check against them is testing something real.
"""

from __future__ import annotations

import json
import sqlite3
from dataclasses import FrozenInstanceError

import pytest

from gaf.checks.contracts import (
    CHECK_IDS,
    FIT_VERDICTS,
    SCOPES,
    CheckFinding,
    CheckReport,
    Severity,
)
from gaf.config import (
    DEFAULT_QUESTION_VARIANT,
    MOCK_REGISTRY,
    QUESTION_VARIANTS,
    CodingRules,
    RunConfig,
)
from gaf.embed.protocol import Embedder, Route, code_text, cosine, route_similarity
from gaf.llm.base import (
    FAIL_SAFE_DEFAULTS,
    LLMError,
    MalformedReplyError,
    TaskType,
    fail_safe_for,
    parse_json_object,
    retry_with_jitter,
    strip_fences,
    validate_enum,
)
from gaf.models import (
    OPERATION_TYPES,
    Assignment,
    Candidate,
    Code,
    Codebook,
    Evidence,
    Operation,
    Response,
    family_of,
    split_name,
)
from gaf.store.schema import SCHEMA_VERSION, TABLES, apply_schema
from gaf.textnorm import NORMALISATION_VERSION, normalise, normalise_for_match
from tests.fixtures import (
    CASES,
    FABRICATED_QUOTE,
    StubEmbedder,
    broken_codebook,
    corpus_by_id,
    human_assignments,
    machine_assignments,
    near_duplicate_codebook,
    response_ids,
    score_counts,
    synthetic_corpus,
    synthetic_scored_table,
    toy_codebook,
    write_coded_xlsx,
    write_narrative_state_xlsx,
)
from tests.fixtures.candidates import word_windows
from tests.fixtures.codebooks import MISSING_RESPONSE_ID, ORPHAN_PARENT_ID

# --------------------------------------------------------------------------- #
# gaf.textnorm — the single normalisation function
# --------------------------------------------------------------------------- #


def test_normalise_is_idempotent_over_the_whole_corpus() -> None:
    for response in synthetic_corpus():
        once = normalise(response.content)
        assert normalise(once) == once


def test_normalise_folds_smart_punctuation_and_collapses_whitespace() -> None:
    assert normalise("don’t   stop\n\nnow") == "don't stop now"
    assert normalise("a — b – c") == "a - b - c"
    assert normalise("  “quoted”  ") == '"quoted"'


def test_normalise_preserves_case_and_match_form_does_not() -> None:
    assert normalise("AI Futures") == "AI Futures"
    assert normalise_for_match("AI Futures") == "ai futures"


def test_normalisation_version_is_declared() -> None:
    assert NORMALISATION_VERSION


# --------------------------------------------------------------------------- #
# gaf.models
# --------------------------------------------------------------------------- #


def test_dataclasses_are_frozen() -> None:
    """Frozen-ness is load-bearing: a checker that wanted to mutate state cannot."""
    response = synthetic_corpus()[0]
    with pytest.raises(FrozenInstanceError):
        response.id = 1  # type: ignore[misc]
    code = next(iter(toy_codebook().codes.values()))
    with pytest.raises(FrozenInstanceError):
        code.name = "x"  # type: ignore[misc]


def test_name_splits_on_the_first_hyphen_only() -> None:
    assert split_name("positive_impacts-problem-solving") == ("positive_impacts", "problem-solving")
    assert split_name("future") == ("future", "")
    assert family_of("AI-superintelligence") == "AI"


def test_codebook_accessors_are_sorted_and_structural() -> None:
    book = toy_codebook()
    assert book.names() == sorted(book.names())
    assert list(book.families()) == sorted(book.families())
    leaf_names = {c.name for c in book.leaves()}
    assert "positive_impacts-healthcare" in leaf_names
    assert "positive_impacts" not in leaf_names  # it has children
    assert book.hierarchy_skeleton()["negative_impacts"] == ["job_destruction", "misuse"]


def test_codebook_json_round_trip_is_lossless_and_deterministic() -> None:
    book = toy_codebook()
    once = book.to_json_str()
    assert Codebook.from_json(json.loads(once)).to_json_str() == once
    # No wall-clock anywhere: that is what makes two runs byte-identical.
    assert "created_at" not in once and "updated_at" not in once


def test_evidence_loader_tolerates_str_and_int_keys() -> None:
    """JSON round-trips stringify integer keys; both shapes must load identically."""
    as_int = Code.from_json({"id": "c", "name": "a-b", "evidence": {9: ["q1", "q2"], 10: ["q3"]}})
    as_str = Code.from_json({"id": "c", "name": "a-b", "evidence": {"9": ["q1", "q2"], "10": ["q3"]}})
    assert [e.to_json() for e in as_int.evidence] == [e.to_json() for e in as_str.evidence]
    assert [e.response_id for e in as_str.evidence] == [9, 9, 10]
    # ...and the canonical list shape round-trips too.
    canonical = Code.from_json(as_str.to_json())
    assert [e.quote for e in canonical.evidence] == ["q1", "q2", "q3"]


def test_bare_string_evidence_is_rejected_not_guessed() -> None:
    with pytest.raises(ValueError, match="bare str"):
        Code.from_json({"id": "c", "name": "a-b", "evidence": ["a quote with no response"]})


def test_candidate_verified_evidence_filters() -> None:
    candidate = Candidate(
        name="a-b",
        description="d",
        evidence=[
            Evidence(203, "yes", verified=True, score=1.0),
            Evidence(203, "no", verified=False, score=0.1),
        ],
    )
    assert [e.quote for e in candidate.verified_evidence()] == ["yes"]


def test_operation_type_whitelist_is_enforced_on_load() -> None:
    assert set(OPERATION_TYPES) == {"create", "merge", "split", "reparent", "rename", "noop"}
    assert Operation.from_json({"type": "split", "targets": ["c1"]}).type == "split"
    with pytest.raises(ValueError, match="unknown operation type"):
        Operation.from_json({"type": "obliterate"})


def test_split_and_reparent_exist_because_the_narrow_op_set_flattened_the_codebook() -> None:
    """Regression guard on the documented cause of codebook flattening (Ng & Chan 2026)."""
    assert "split" in OPERATION_TYPES
    assert "reparent" in OPERATION_TYPES


def test_response_meta_round_trips_and_is_separate_from_content() -> None:
    response = Response(id=1, question="q", content="c", source="s", meta={"age": 30})
    assert Response.from_json(response.to_json()).meta == {"age": 30}


def test_assignment_round_trip() -> None:
    row = Assignment(9, "a segment", "a-b")
    assert Assignment.from_json(row.to_json()) == row


# --------------------------------------------------------------------------- #
# gaf.checks.contracts
# --------------------------------------------------------------------------- #


def test_check_report_severity_semantics() -> None:
    report = CheckReport()
    report.add("S1", Severity.WARN, "candidate", "a-b", "no description")
    report.add("S2", Severity.INFO, "candidate", "a-b", "fuzzy match", score=0.91)
    assert report.passed()
    report.add("S2", Severity.ERROR, "candidate", "c-d", "no verified evidence")
    assert not report.passed()
    assert len(report.errors()) == 1
    assert len(report.warnings()) == 1
    assert len(report.infos()) == 1


def test_check_report_summary_and_totals_are_deterministic() -> None:
    report = CheckReport()
    report.add("S2", Severity.INFO, "candidate", "x", "m")
    report.add("S1", Severity.WARN, "candidate", "y", "m")
    report.add("S1", Severity.WARN, "candidate", "z", "m")
    assert report.summary() == {"S1": {"WARN": 2}, "S2": {"INFO": 1}}
    assert list(report.summary()) == ["S1", "S2"]  # sorted
    assert report.totals() == {"ERROR": 0, "WARN": 2, "INFO": 1}


def test_check_report_extend_preserves_order() -> None:
    first, second = CheckReport(), CheckReport()
    first.add("S1", Severity.INFO, "candidate", "1", "a")
    second.add("S2", Severity.INFO, "candidate", "2", "b")
    first.extend(second)
    assert [f.subject for f in first] == ["1", "2"]


def test_check_finding_json_round_trip() -> None:
    finding = CheckFinding("M3", Severity.WARN, "candidate", "a-b", "low fit", {"fit": 0.21})
    assert CheckFinding.from_json(finding.to_json()) == finding


def test_check_ids_and_verdicts_are_the_documented_sets() -> None:
    assert CHECK_IDS == ("S1", "S2", "S2b", "S3", "S4", "S5", "S6", "M1", "M2", "M3", "M4")
    assert FIT_VERDICTS == ("APPLIES", "IMPRECISE", "INCOMPLETE", "UNNECESSARY")
    assert "segment" in SCOPES and "pair" in SCOPES


def test_check_report_has_no_handle_on_state() -> None:
    """Structural guarantee that a checker cannot decide anything: it holds only findings."""
    assert set(CheckReport.__dataclass_fields__) == {"findings"}


# --------------------------------------------------------------------------- #
# gaf.config
# --------------------------------------------------------------------------- #


def test_coding_rules_match_the_pi_instructions() -> None:
    rules = CodingRules()
    assert (rules.min_codes_per_response, rules.max_codes_per_response) == (2, 12)
    assert rules.max_codes_per_segment == 2
    assert rules.max_quote_sentences == 2
    assert rules.hierarchy_depth == 2
    assert rules.prefer_subcode_first is True
    assert (rules.tau_high, rules.tau_low, rules.tau_fit) == (0.80, 0.45, 0.30)
    assert rules.fuzzy_threshold == 0.85
    assert rules.segment_overlap_threshold == 0.5


def test_run_config_defaults_are_offline_and_seeded() -> None:
    config = RunConfig()
    assert config.offline is True
    assert config.models is MOCK_REGISTRY
    assert config.embedding.mode == "lexical"
    assert config.recycle_human_corrections is False  # held out, per brief §16.6
    assert config.checkpoints.hard_floor_responses == 50
    assert config.analysis.min_code_frequency == 2
    assert isinstance(config.seed, int)


def test_run_config_json_is_complete_and_serialisable() -> None:
    payload = RunConfig().to_json()
    json.dumps(payload)  # must not raise: it is written into the runs table
    for key in ("rules", "models", "embedding", "checkpoints", "analysis", "lexical"):
        assert key in payload


def test_both_question_variants_are_present_and_the_default_is_declared() -> None:
    assert set(QUESTION_VARIANTS) == {"v2", "v3"}
    assert DEFAULT_QUESTION_VARIANT == "v2"
    assert QUESTION_VARIANTS["v2"].startswith("From now to 2050")
    assert QUESTION_VARIANTS["v3"].startswith("In 2050, what kind of roles")
    assert RunConfig().resolved_question() == QUESTION_VARIANTS["v2"]


def test_mock_registry_is_free_and_the_live_coders_come_from_different_providers() -> None:
    from gaf.config import DEFAULT_LIVE_REGISTRY

    assert all(spec.provider == "mock" for spec in MOCK_REGISTRY.all_specs())
    assert MOCK_REGISTRY.coder_a.cost_usd(1_000, 1_000) == 0.0
    assert DEFAULT_LIVE_REGISTRY.distinct_coder_providers()
    assert DEFAULT_LIVE_REGISTRY.judge.provider not in {
        DEFAULT_LIVE_REGISTRY.coder_a.provider,
        DEFAULT_LIVE_REGISTRY.coder_b.provider,
    }


def test_every_registry_entry_names_a_model_id_its_provider_could_serve() -> None:
    """A provider and a model id that disagree is a live run that dies on its first call.

    The prefixes are the providers' own naming conventions, not this package's: OpenAI
    ships `gpt-*`, Google `gemini-*`, Anthropic `claude-*`, and the deterministic
    stand-ins in `gaf.llm.mock` answer to `mock-*` (ADR-0048).
    """
    from gaf.config import DEFAULT_LIVE_REGISTRY

    prefixes = {
        "mock": "mock-",
        "openai": "gpt-",
        "gemini": "gemini-",
        "anthropic": "claude-",
    }
    for registry in (MOCK_REGISTRY, DEFAULT_LIVE_REGISTRY):
        for spec in registry.all_specs():
            assert spec.model.startswith(prefixes[spec.provider]), (spec.role, spec.model)
    # The epistemic-diversity mechanism, not a preference: it must survive a re-binding.
    assert DEFAULT_LIVE_REGISTRY.distinct_coder_providers()


def test_the_live_registry_carries_the_prices_adr_0048_read_and_the_date_it_read_them() -> None:
    """Pinned to ADR-0048, so a silent drift fails here rather than in a run report.

    `ModelSpec.cost_usd` is the only place a rate lives, so these four pairs are what
    `stats.llm.cost_usd` would print on a live run. Moving one is a decision to record,
    not an edit to make: change the ADR and this test together, or neither.
    """
    from gaf.config import DEFAULT_LIVE_REGISTRY

    priced = {
        (spec.role, spec.provider, spec.model): (
            spec.input_usd_per_mtok,
            spec.output_usd_per_mtok,
        )
        for spec in DEFAULT_LIVE_REGISTRY.all_specs()
    }
    assert priced == {
        ("coder_a", "openai", "gpt-4o-mini"): (0.15, 0.60),
        ("coder_b", "gemini", "gemini-3.6-flash"): (0.75, 3.75),
        ("judge", "anthropic", "claude-sonnet-5"): (2.00, 10.00),
        ("refactorer", "anthropic", "claude-opus-5"): (5.00, 25.00),
    }
    assert all(spec.priced_on == "2026-09-23" for spec in DEFAULT_LIVE_REGISTRY.all_specs())


def test_a_price_date_travels_with_a_live_spec_and_never_with_a_mock_one() -> None:
    """`priced_on` is omitted from the JSON when empty, which is what keeps the offline
    artefacts byte-identical: the golden manifest and the committed example both carry
    `RunConfig.to_json()` in full, and a mock has no price to date (ADR-0048)."""
    from gaf.config import DEFAULT_LIVE_REGISTRY

    for spec in MOCK_REGISTRY.all_specs():
        assert spec.priced_on == ""
        assert "priced_on" not in spec.to_json()
    for spec in DEFAULT_LIVE_REGISTRY.all_specs():
        assert spec.to_json()["priced_on"] == "2026-09-23"
    offline = RunConfig().to_json()["models"]
    assert all("priced_on" not in spec for spec in offline.values())


def test_model_registry_role_lookup() -> None:
    assert MOCK_REGISTRY.for_role("judge") is MOCK_REGISTRY.judge
    with pytest.raises(ValueError):
        MOCK_REGISTRY.for_role("slicer")  # there is no slicer agent, by design


# --------------------------------------------------------------------------- #
# gaf.embed.protocol
# --------------------------------------------------------------------------- #


def test_stub_embedder_satisfies_the_protocol_and_is_deterministic() -> None:
    embedder = StubEmbedder()
    assert isinstance(embedder, Embedder)
    first = embedder.embed_one("machines will take the jobs")
    second = StubEmbedder().embed_one("machines will take the jobs")
    assert (first == second).all()
    assert abs(float((first**2).sum()) - 1.0) < 1e-9  # L2-normalised
    assert embedder.embed([]).shape == (0, embedder.dim)
    assert embedder.embed(["a", "b"]).shape == (2, embedder.dim)


def test_code_text_is_the_one_rendering() -> None:
    assert code_text("a-b", "desc") == "a-b: desc"
    assert code_text("a-b", "") == "a-b"
    assert code_text("  a-b  ", "  desc  ") == "a-b: desc"


def test_routing_bands() -> None:
    rules = CodingRules()
    assert route_similarity(0.95, rules.tau_high, rules.tau_low) is Route.MERGE
    assert route_similarity(0.80, rules.tau_high, rules.tau_low) is Route.MERGE  # inclusive
    assert route_similarity(0.60, rules.tau_high, rules.tau_low) is Route.JUDGE
    assert route_similarity(0.45, rules.tau_high, rules.tau_low) is Route.JUDGE  # inclusive
    assert route_similarity(0.10, rules.tau_high, rules.tau_low) is Route.CREATE


def test_cosine_is_safe_on_zero_vectors() -> None:
    embedder = StubEmbedder()
    assert cosine(embedder.embed_one(""), embedder.embed_one("anything")) == 0.0


# --------------------------------------------------------------------------- #
# gaf.llm.base
# --------------------------------------------------------------------------- #


def test_fence_stripping_is_idempotent() -> None:
    assert strip_fences('```json\n{"a": 1}\n```') == '{"a": 1}'
    assert strip_fences('{"a": 1}') == '{"a": 1}'
    assert strip_fences(strip_fences("```\n{}\n```")) == "{}"


def test_json_parsing_rescues_chatty_replies_and_refuses_junk() -> None:
    assert parse_json_object('Sure! {"verdict": "APPLIES"} Hope that helps.') == {
        "verdict": "APPLIES"
    }
    with pytest.raises(MalformedReplyError):
        parse_json_object("I would rather not.")
    with pytest.raises(MalformedReplyError):
        parse_json_object("[1, 2, 3]")  # an array is not an object


def test_enum_whitelisting_tolerates_case_and_rejects_invention() -> None:
    assert validate_enum(" applies ", FIT_VERDICTS, "APPLIES") == "APPLIES"
    assert validate_enum("MOSTLY_FINE", FIT_VERDICTS, "APPLIES") == "APPLIES"
    assert validate_enum(None, FIT_VERDICTS, "APPLIES") == "APPLIES"
    assert validate_enum("UNNECESSARY", FIT_VERDICTS, "APPLIES") == "UNNECESSARY"


def test_every_task_type_has_a_non_destructive_fail_safe() -> None:
    for task in TaskType:
        assert task in FAIL_SAFE_DEFAULTS, f"{task} has no fail-safe default"
    # The specific non-destructive choices, each of which protects data:
    assert FAIL_SAFE_DEFAULTS[TaskType.JUDGE_FIT]["verdict"] == "APPLIES"  # keeps the quote
    assert FAIL_SAFE_DEFAULTS[TaskType.JUDGE_ROUTE]["route"] == "CREATE"  # never a blind merge
    assert FAIL_SAFE_DEFAULTS[TaskType.CODE]["candidates"] == []  # invents nothing
    assert FAIL_SAFE_DEFAULTS[TaskType.REFACTOR]["operations"] == []  # a no-op edit script


def test_fail_safe_returns_a_copy_not_the_shared_default() -> None:
    first = fail_safe_for(TaskType.JUDGE_FIT)
    first["verdict"] = "UNNECESSARY"
    assert FAIL_SAFE_DEFAULTS[TaskType.JUDGE_FIT]["verdict"] == "APPLIES"


def test_retry_is_seeded_and_eventually_raises() -> None:
    attempts: list[int] = []

    def flaky() -> str:
        attempts.append(1)
        if len(attempts) < 3:
            raise TimeoutError("transient")
        return "ok"

    result, used = retry_with_jitter(flaky, attempts=3, seed=7, sleep=lambda _: None)
    assert (result, used) == ("ok", 3)

    with pytest.raises(LLMError):
        retry_with_jitter(
            lambda: (_ for _ in ()).throw(TimeoutError("always")),
            attempts=2,
            seed=7,
            sleep=lambda _: None,
        )


# --------------------------------------------------------------------------- #
# gaf.store.schema
# --------------------------------------------------------------------------- #


def test_schema_applies_and_is_idempotent() -> None:
    conn = sqlite3.connect(":memory:")
    apply_schema(conn)
    apply_schema(conn)  # opening an existing database is the same call
    present = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert set(TABLES) <= present
    version = conn.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]
    assert version == str(SCHEMA_VERSION)


@pytest.fixture()
def db() -> sqlite3.Connection:
    conn = sqlite3.connect(":memory:")
    apply_schema(conn)
    conn.execute("INSERT INTO runs VALUES ('r1', '2026-09-01T00:00:00Z', '0.1.0', 1, '{}')")
    return conn


def test_audit_log_is_append_only(db: sqlite3.Connection) -> None:
    db.execute("INSERT INTO audit(run_id, created_at, event) VALUES ('r1', 't', 'coded')")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        db.execute("UPDATE audit SET event = 'tampered'")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        db.execute("DELETE FROM audit")
    assert db.execute("SELECT COUNT(*) FROM audit").fetchone()[0] == 1


def test_snapshots_are_immutable(db: sqlite3.Connection) -> None:
    db.execute(
        "INSERT INTO snapshots(snapshot_id, run_id, created_at, reason, code_count, codebook_json)"
        " VALUES ('s1', 'r1', 't', 'seed', 0, '{}')"
    )
    with pytest.raises(sqlite3.IntegrityError, match="immutable"):
        db.execute("UPDATE snapshots SET codebook_json = '{\"codes\": []}'")
    with pytest.raises(sqlite3.IntegrityError, match="immutable"):
        db.execute("DELETE FROM snapshots")


def test_foreign_keys_are_enforced(db: sqlite3.Connection) -> None:
    with pytest.raises(sqlite3.IntegrityError):
        db.execute(
            "INSERT INTO codings(coding_id, run_id, response_id, source, coder, snapshot_id,"
            " prompt_version, created_at, raw_json)"
            " VALUES ('c1', 'nonexistent-run', 3, 'synthetic', 'coder_a', 'nope', 'v1', 't', '{}')"
        )


# --------------------------------------------------------------------------- #
# Fixtures actually contain what they claim
# --------------------------------------------------------------------------- #


def test_corpus_shape_and_non_contiguous_ids() -> None:
    corpus = synthetic_corpus()
    assert len(corpus) >= 12
    ids = response_ids()
    assert ids == sorted(ids) and len(set(ids)) == len(ids)
    assert ids != list(range(len(ids)))  # non-contiguous, like the real data
    assert all(r.question and r.content and r.source for r in corpus)


def test_every_planted_quote_is_locatable_except_the_fabricated_one() -> None:
    corpus = corpus_by_id()
    for name, planted in CASES.items():
        for candidate in planted.candidates:
            for evidence in candidate.evidence:
                if evidence.quote == FABRICATED_QUOTE:
                    continue
                text = normalise(corpus[evidence.response_id].content)
                assert evidence.quote in text, f"{name}: {candidate.name} quote not in source"


def test_fabricated_quote_appears_in_no_response() -> None:
    for response in synthetic_corpus():
        assert FABRICATED_QUOTE not in normalise(response.content)


def test_a_case_exists_for_every_check() -> None:
    covered = {planted.check_id for planted in CASES.values()}
    for check_id in ("S1", "S2", "S2b", "S3", "S4", "S5", "M3"):
        assert check_id in covered, f"no planted case for {check_id}"


def test_s4_case_really_puts_three_codes_on_one_identical_span() -> None:
    planted = CASES["S4_three_codes_one_segment"]
    quotes = {e.quote for c in planted.candidates for e in c.evidence}
    assert len(planted.candidates) == 3
    assert len(quotes) == 1  # the same segment, verbatim


def test_s5_too_many_case_uses_non_overlapping_windows() -> None:
    """Otherwise it would trip S4 as well and the test would prove nothing."""
    windows = word_windows(218, 13, width=4)
    assert len(windows) == len(set(windows)) == 13


def test_broken_codebook_contains_every_s6_violation() -> None:
    book = broken_codebook()
    names = [c.name.casefold() for c in book.codes.values()]
    assert len(names) != len(set(names))  # duplicate, case-insensitively
    assert any(c.parent_id == ORPHAN_PARENT_ID for c in book.codes.values())
    assert any(c.description == "" for c in book.codes.values())
    assert any(
        e.response_id == MISSING_RESPONSE_ID for c in book.codes.values() for e in c.evidence
    )
    deep = book.codes["c-depth-deep"]
    mid = book.codes[str(deep.parent_id)]
    assert mid.parent_id == "c-depth-root"  # a three-deep chain
    parent_with_evidence = book.codes["c-parent-with-evidence"]
    assert parent_with_evidence.evidence and book.children_of(parent_with_evidence.id)


def test_near_duplicate_pair_is_clear_of_both_thresholds() -> None:
    rules = CodingRules()
    book = near_duplicate_codebook()
    embedder = StubEmbedder()

    def similarity(a: str, b: str) -> float:
        first, second = book.by_name(a), book.by_name(b)
        assert first and second
        return cosine(
            embedder.embed_one(code_text(first.name, first.description)),
            embedder.embed_one(code_text(second.name, second.description)),
        )

    duplicate = similarity("negative_impacts-job_destruction", "negative_impacts-job_loss")
    control = similarity("negative_impacts-job_destruction", "negative_impacts-dependence")
    assert duplicate > rules.tau_high + 0.02, f"near-duplicate pair too close to tau_high: {duplicate}"
    assert control < rules.tau_low - 0.02, f"control pair too close to tau_low: {control}"


def test_scored_table_is_deterministic_and_fittable() -> None:
    rows = synthetic_scored_table()
    assert [r.text for r in synthetic_scored_table()] == [r.text for r in rows]
    counts = score_counts(rows)
    assert counts["n"] == 240
    # Both binarised cuts must clear the refuse-to-fit floor of 10 per class.
    for cut in ("ge1", "ge3"):
        assert counts[cut]["positive"] >= 10 and counts[cut]["negative"] >= 10
    # score == 2 must exist in quantity: ge3 counts it as a NEGATIVE, never a dropped row.
    assert counts["score_2_rows"] >= 20
    assert sum(1 for r in rows if r.score == 2 and r.score < 3) == counts["score_2_rows"]


def test_assignment_fixtures_have_both_unmatched_directions() -> None:
    """The over-coding and blind-spot lists are the write-up's most valuable output."""
    human = {a.code for a in human_assignments()}
    machine = {a.code for a in machine_assignments()}
    assert machine - human, "no machine-only codes: over-coding would be untestable"
    assert human - machine, "no human-only codes: blind spots would be untestable"


def test_xlsx_builders_write_both_header_shapes(tmp_path) -> None:  # type: ignore[no-untyped-def]
    from openpyxl import load_workbook

    from tests.fixtures.xlsx import CODED_HEADER_VARIANTS

    raw = write_narrative_state_xlsx(tmp_path / "ns.xlsx")
    rows = list(load_workbook(raw).active.iter_rows(values_only=True))
    assert rows[0] == ("Number", "Response")
    assert [r[0] for r in rows[1:]] == [9, 10, 12]  # non-contiguous

    for header in CODED_HEADER_VARIANTS:
        coded = write_coded_xlsx(
            tmp_path / f"coded_{header[0]}.xlsx",
            [(9, "a segment", "positive_impacts-healthcare")],
            header=header,
        )
        assert next(load_workbook(coded).active.iter_rows(values_only=True)) == header


def test_no_real_survey_data_is_reachable_from_the_fixtures() -> None:
    """Every fixture response is invented; none is read from disk."""
    from tests.fixtures import corpus as corpus_module

    source = (corpus_module.__file__ or "")
    assert source.endswith("corpus.py")
    assert all(r.source == "synthetic" for r in synthetic_corpus())
